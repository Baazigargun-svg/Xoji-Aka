import sqlite3
import telebot
import os
import threading
from flask import Flask, render_template_string, request
from telebot import types
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BOT_TOKEN = "8573337094:AAHH1kNQnNrJyMfG6d5z6IO8lgkS1UdurR8"
ADMIN_ID = 6851851908  
SEX_GROUP_ID = -6851851908 

bot = telebot.TeleBot(BOT_TOKEN)
app = Flask(__name__)
user_steps = {}

def init_db():
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute('''CREATE TABLE IF NOT EXISTS users (tg_id INTEGER PRIMARY KEY, name TEXT, phone TEXT, role TEXT DEFAULT 'pending')''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS products (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE, optom_price REAL, chakana_price REAL, stock REAL DEFAULT 0, category TEXT DEFAULT 'Boshqa', cost_price REAL DEFAULT 0)''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS shops (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE, phone TEXT, debt REAL DEFAULT 0)''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS orders (id INTEGER PRIMARY KEY AUTOINCREMENT, shop_name TEXT, agent_name TEXT, total_sum REAL, items_text TEXT, price_type TEXT, status TEXT DEFAULT 'Yangi', date TEXT)''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS incomes (id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT, amount REAL, date TEXT)''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS expenses (id INTEGER PRIMARY KEY AUTOINCREMENT, reason TEXT, amount REAL, date TEXT)''')
    
    try:
        cursor.execute("ALTER TABLE products ADD COLUMN chakana_price REAL;")
    except sqlite3.OperationalError:
        pass

    conn.commit()
    conn.close()

def auto_insert_products():
    products_list = [
        ("Pelmen 300 gr", 14000), ("Pelmen 500 gr", 24000), ("Pelmen frikadel 500 gr", 28000),
        ("Teftel 300 gr", 23000), ("Golubtsi 500 gr", 25000), ("Tok do'lma 300 gr", 25000),
        ("Karam do'lma 300 gr", 25000), ("Somsa kesilgani (20 dona)", 17000), ("O'rama hamir", 17000),
        ("Osh masalliq 0.5 kg", 13000), ("Osh masalliq 1 kg", 15000), ("Farsh mol go'shtidan", 19000),
        ("Frikadel 300 gr", 20000), ("Frikadel 500 gr", 30000), ("KFC 500 gr", 30000),
        ("Mantini hamiri 25 dona", 13000), ("Kotlet 5 dona", 25000)
    ]
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    for name, price in products_list:
        try:
            cursor.execute("INSERT INTO products (name, optom_price, chakana_price, stock, category) VALUES (?, ?, ?, 0, 'Yarim tayyor')", (name, price, price))
        except sqlite3.IntegrityError:
            continue
    conn.commit()
    conn.close()

def create_excel_invoice(order_id, shop_name, agent_name, date_str, price_type, cart_items):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Nakladnoy_{order_id}"
    ws.sheet_view.showGridLines = True
    
    title_font = Font(name='Arial', size=16, bold=True)
    header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    bold_font = Font(name='Arial', size=11, bold=True)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    total_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    thin_border = Border(left=Side(style='thin', color='B0B0B0'), right=Side(style='thin', color='B0B0B0'), top=Side(style='thin', color='B0B0B0'), bottom=Side(style='thin', color='B0B0B0'))
    
    ws.merge_cells('A1:E1')
    ws['A1'] = "XOJI AKA FACTORY — NAKLADNOY"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')
    
    ws['A3'] = f"Buyurtma ID: #{order_id}"
    ws['A3'].font = bold_font
    ws['D3'] = f"Sana: {date_str}"
    ws['A4'] = f"Do'kon (Klient): {shop_name}"
    ws['D4'] = f"Narx turi: {price_type.upper()}"
    ws['A5'] = f"Agent: {agent_name}"
    
    headers = ["№", "Mahsulot nomi", "Miqdori", "Narxi (so'm)", "Jami summa"]
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=7, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
        
    row_num = 8
    total_sum = 0
    for idx, item in enumerate(cart_items, 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=item['name']).alignment = Alignment(horizontal='left')
        ws.cell(row=row_num, column=3, value=item['qty']).alignment = Alignment(horizontal='right')
        ws.cell(row=row_num, column=4, value=item['price']).alignment = Alignment(horizontal='right')
        summa = item['qty'] * item['price']
        total_sum += summa
        ws.cell(row=row_num, column=5, value=summa).alignment = Alignment(horizontal='right')
        ws.cell(row=row_num, column=4).number_format = '#,##0'
        ws.cell(row=row_num, column=5).number_format = '#,##0'
        for col in range(1, 6): ws.cell(row=row_num, column=col).border = thin_border
        row_num += 1
        
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
    ws.cell(row=row_num, column=1, value="JAMI TO'LOV:").alignment = Alignment(horizontal='right')
    ws.cell(row=row_num, column=1).font = bold_font
    total_val = ws.cell(row=row_num, column=5, value=total_sum)
    total_val.font = bold_font
    total_val.number_format = '#,##0'
    for col in range(1, 6):
        ws.cell(row=row_num, column=col).fill = total_fill
        ws.cell(row=row_num, column=col).border = thin_border
        
    row_num += 2
    ws.cell(row=row_num, column=2, value="Qabul qildim: ____")
    ws.cell(row=row_num, column=4, value="Topshirdim: ____")

    for col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col)
        max_len = max(len(str(ws.cell(row=row, column=col).value or '')) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    file_name = f"Nakladnoy_{order_id}.xlsx"
    wb.save(file_name)
    return file_name

def get_main_menu(role):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    if role == "admin":
        markup.row(types.KeyboardButton("📦 Sklad & Mahsulotlar"), types.KeyboardButton("📊 Kunlik Hisobot"))
        markup.row(types.KeyboardButton("👥 Agentlar boshqaruvi"), types.KeyboardButton("🏪 AKB (Do'konlar & Qarz)"))
    elif role == "agent":
        markup.row(types.KeyboardButton("🛒 Yangi Buyurtma Urish"), types.KeyboardButton("🏪 Do'kon qo'shish (AKB)"))
        markup.row(types.KeyboardButton("💰 Qarz/To'lov yozish"), types.KeyboardButton("📜 Mening Buyurtmalarim"))
    else:
        markup.add(types.KeyboardButton("📝 Ro'yxatdan o'tish"))
    return markup

@bot.message_handler(commands=['start'])
def start_command(message):
    tg_id = message.from_user.id
    init_db()
    auto_insert_products()
    if tg_id == ADMIN_ID:
        conn = sqlite3.connect("xoji_aka_factory.db")
        cursor = conn.cursor()
        cursor.execute("INSERT OR REPLACE INTO users (tg_id, name, phone, role) VALUES (?, 'Admin', '', 'admin')", (tg_id,))
        conn.commit()
        conn.close()
        bot.send_message(message.chat.id, "Xoji aka, xush kelibsiz! Boshqaruv paneli tayyor.", reply_markup=get_main_menu("admin"))
        return
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT role, name FROM users WHERE tg_id = ?", (tg_id,))
    user = cursor.fetchone()
    conn.close()
    if user:
        role, name = user[0], user[1]
        if role == "pending":
            bot.send_message(message.chat.id, f"Salom {name}. So'rovingiz admin tasdig'ini kutyapti.")
        else:
            bot.send_message(message.chat.id, f"Salom {name}! Ishni boshlashimiz mumkin.", reply_markup=get_main_menu(role))
    else:
        bot.send_message(message.chat.id, "Assalomu alaykum! Tizimga xush kelibsiz. Davom etish uchun ro'yxatdan o'ting.", reply_markup=get_main_menu("guest"))

@bot.message_handler(func=lambda message: message.text == "📦 Sklad & Mahsulotlar")
def admin_sklad_menu(message):
    if message.from_user.id != ADMIN_ID: return
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT id, name, stock FROM products")
    prods = cursor.fetchall()
    conn.close()

    inline_kb = types.InlineKeyboardMarkup(row_width=1)
    if not prods:
        bot.send_message(message.chat.id, "📦 Omborxonada mahsulotlar qolmagan.")
        return
    for p in prods:
        stock_val = p[2] if p[2] is not None else 0
        inline_kb.add(types.InlineKeyboardButton(f"🔹 {p[1]} ({int(stock_val)} kg/dona)", callback_data=f"adm_prod_{p[0]}"))
    inline_kb.add(types.InlineKeyboardButton("🆕 ✨ YANGI MAHSULOT QO'SHISH", callback_data="adm_create_product"))
    bot.send_message(message.chat.id, "<b>📦 Sklad nazorati:</b>", parse_mode="HTML", reply_markup=inline_kb)

@bot.callback_query_handler(func=lambda call: call.data.startswith("adm_prod_"))
def admin_product_detail(call):
    p_id = int(call.data.split("_")[-1])
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT name, optom_price, chakana_price, stock FROM products WHERE id = ?", (p_id,))
    p = cursor.fetchone()
    conn.close()
    if p:
        optom = p[1] if p[1] is not None else 0
        chakana = p[2] if p[2] is not None else 0
        stock = p[3] if p[3] is not None else 0
        text = f"📦 <b>Mahsulot:</b> {p[0]}\n💰 Optom: {optom:,.0f} so'm\n🛍 Chakana: {chakana:,.0f} so'm\n🔢 <b>Qoldiq:</b> {int(stock)}"
        markup = types.InlineKeyboardMarkup()
        markup.row(types.InlineKeyboardButton("➕ Qoldiq Qo'shish", callback_data=f"stk_plus_{p_id}"), types.InlineKeyboardButton("➖ Qoldiq Ayirish", callback_data=f"stk_minus_{p_id}"))
        markup.row(types.InlineKeyboardButton("💵 Optom Narx", callback_data=f"prc_optom_{p_id}"), types.InlineKeyboardButton("💵 Chakana Narx", callback_data=f"prc_chakana_{p_id}"))
        markup.row(types.InlineKeyboardButton("🗑 O'chirish", callback_data=f"stk_del_{p_id}"))
        bot.edit_message_text(text, call.message.chat.id, call.message.message_id, parse_mode="HTML", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith(("stk_", "prc_")))
def admin_stock_price_actions(call):
    prefix, action, p_id = call.data.split("_")
    p_id = int(p_id)
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM products WHERE id = ?", (p_id,))
    res = cursor.fetchone()
    p_name = res[0] if res else "Nomalum"
    conn.close()

    if action == "del":
        conn = sqlite3.connect("xoji_aka_factory.db")
        cursor = conn.cursor()
        cursor.execute("DELETE FROM products WHERE id = ?", (p_id,))
        conn.commit()
        conn.close()
        bot.send_message(call.message.chat.id, f"🗑 <b>{p_name}</b> o'chirildi.", parse_mode="HTML")
        return

    msg = bot.send_message(call.message.chat.id, f"🔢 <b>{p_name}</b> uchun qiymat kiriting:", parse_mode="HTML")
    bot.register_next_step_handler(msg, save_product_edits, prefix, action, p_id)

def save_product_edits(message, prefix, action, p_id):
    try:
        val = float(message.text)
        conn = sqlite3.connect("xoji_aka_factory.db")
        cursor = conn.cursor()
        if action == "plus": cursor.execute("UPDATE products SET stock = stock + ? WHERE id = ?", (val, p_id))
        elif action == "minus": cursor.execute("UPDATE products SET stock = stock - ? WHERE id = ?", (val, p_id))
        elif action == "optom": cursor.execute("UPDATE products SET optom_price = ? WHERE id = ?", (val, p_id))
        elif action == "chakana": cursor.execute("UPDATE products SET chakana_price = ? WHERE id = ?", (val, p_id))
        conn.commit()
        conn.close()
        bot.send_message(message.chat.id, "✅ O'zgarish saqlandi!")
    except: bot.send_message(message.chat.id, "❌ Faqat raqam kiriting.")

@bot.callback_query_handler(func=lambda call: call.data == "adm_create_product")
def admin_create_product_start(call):
    msg = bot.send_message(call.message.chat.id, "📝 Yangi mahsulot NOMINI kiriting:")
    bot.register_next_step_handler(msg, process_new_p_name)

def process_new_p_name(message):
    name = message.text
    msg = bot.send_message(message.chat.id, f"💰 '{name}' uchun Optom narxini kiriting:")
    bot.register_next_step_handler(msg, process_new_p_optom, name)

def process_new_p_optom(message, name):
    try:
        optom = float(message.text)
        msg = bot.send_message(message.chat.id, f"🛍 '{name}' uchun Chakana narxini kiriting:")
        bot.register_next_step_handler(msg, process_new_p_final, name, optom)
    except: bot.send_message(message.chat.id, "Xato! Faqat raqam kiriting.")

def process_new_p_final(message, name, optom):
    try:
        chakana = float(message.text)
        conn = sqlite3.connect("xoji_aka_factory.db")
        cursor = conn.cursor()
        cursor.execute("INSERT INTO products (name, optom_price, chakana_price, stock) VALUES (?, ?, ?, 0)", (name, optom, chakana))
        conn.commit()
        conn.close()
        bot.send_message(message.chat.id, f"✅ Yangi mahsulot qo'shildi: {name}")
    except: bot.send_message(message.chat.id, "Xato! Bu nom allaqachon mavjud.")

@bot.message_handler(func=lambda message: message.text in ["📊 Kunlik Hisobot", "🏪 AKB (Do'konlar & Qarz)", "👥 Agentlar boshqaruvi"])
def admin_other_sections(message):
    if message.from_user.id != ADMIN_ID: return
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    
    if message.text == "📊 Kunlik Hisobot":
        bugun = datetime.now().strftime("%Y-%m-%d")
        cursor.execute("SELECT id, shop_name, total_sum, status FROM orders WHERE date LIKE ?", (f"{bugun}%",))
        orders = cursor.fetchall()
        text = f"📊 <b>Bugungi buyurtmalar ({bugun}):</b>\n\n"
        inline_kb = types.InlineKeyboardMarkup()
        if not orders: text += "Hali buyurtma yo'q."
        for o in orders:
            t_sum = o[2] if o[2] is not None else 0
            text += f"🆔 #{o[0]} | {o[1]} | {t_sum:,.0f} so'm | <b>{o[3]}</b>\n"
            inline_kb.add(types.InlineKeyboardButton(f"⚙️ #{o[0]} Statusi", callback_data=f"mng_ord_{o[0]}"))
        bot.send_message(message.chat.id, text, parse_mode="HTML", reply_markup=inline_kb)
        
    elif message.text == "🏪 AKB (Do'konlar & Qarz)":
        cursor.execute("SELECT name, phone, debt FROM shops")
        shops = cursor.fetchall()
        text = "<b>🏪 Do'konlar qarzlari:</b>\n\n"
        for s in shops: 
            debt_val = s[2] if s[2] is not None else 0
            text += f"🏢 {s[0]} ({s[1]}) — Qarz: {debt_val:,.0f} so'm\n"
        inline_kb = types.InlineKeyboardMarkup().add(types.InlineKeyboardButton("➕ Do'kon Qo'shish", callback_data="admin_add_shop"))
        bot.send_message(message.chat.id, text, parse_mode="HTML", reply_markup=inline_kb)
        
    elif message.text == "👥 Agentlar boshqaruvi":
        cursor.execute("SELECT name, phone, role, tg_id FROM users WHERE role != 'admin'")
        agents = cursor.fetchall()
        inline_kb = types.InlineKeyboardMarkup()
        text = "<b>👥 Agentlar:</b>\n\n"
        for a in agents:
            status = "✅ Faol" if a[2] == 'agent' else "⏳ Kutilmoqda"
            text += f"👤 {a[0]} ({a[1]}) - {status}\n"
            if a[2] == 'pending': inline_kb.add(types.InlineKeyboardButton(f"👍 {a[0]}ni tasdiqlash", callback_data=f"approve_{a[3]}"))
        bot.send_message(message.chat.id, text, parse_mode="HTML", reply_markup=inline_kb)
    conn.close()

@bot.callback_query_handler(func=lambda call: call.data.startswith("mng_ord_"))
def manage_order_status_menu(call):
    order_id = int(call.data.split("_")[-1])
    markup = types.InlineKeyboardMarkup()
    markup.row(types.InlineKeyboardButton("🔄 Yangi", callback_data=f"st_yangi_{order_id}"), types.InlineKeyboardButton("🚚 Otgruzka", callback_data=f"st_otgruzka_{order_id}"))
    markup.row(types.InlineKeyboardButton("✅ Yetkazildi", callback_data=f"st_done_{order_id}"), types.InlineKeyboardButton("❌ Bekor", callback_data=f"st_otmen_{order_id}"))
    bot.send_message(call.message.chat.id, f"🆔 #{order_id} - Statusni tanlang:", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("st_"))
def change_status_logic(call):
    _, mode, order_id = call.data.split("_")
    order_id = int(order_id)
    status_map = {"yangi": "Yangi", "otgruzka": "Otgruzka", "done": "Yetkazildi", "otmen": "Bekor"}
    new_status = status_map[mode]
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("UPDATE orders SET status = ? WHERE id = ?", (new_status, order_id))
    conn.commit()
    conn.close()
    bot.send_message(call.message.chat.id, f"✅ Status: <b>{new_status}</b>", parse_mode="HTML")

@bot.callback_query_handler(func=lambda call: call.data == "admin_add_shop")
def admin_inline_clicks(call):
    msg = bot.send_message(call.message.chat.id, "🏢 Do'kon NOMINI kiriting:")
    bot.register_next_step_handler(msg, process_shop_name_step)

@bot.message_handler(func=lambda message: message.text == "🏪 Do'kon qo'shish (AKB)")
def agent_add_shop_menu(message):
    msg = bot.send_message(message.chat.id, "🏢 Do'kon NOMINI kiriting:")
    bot.register_next_step_handler(msg, process_shop_name_step)

def process_shop_name_step(message):
    shop_name = message.text
    msg = bot.send_message(message.chat.id, f"📞 '{shop_name}' uchun TELEFON RAqam:")
    bot.register_next_step_handler(msg, process_shop_phone_final, shop_name)

def process_shop_phone_final(message, shop_name):
    phone = message.text
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO shops (name, phone, debt) VALUES (?, ?, 0)", (shop_name, phone))
        conn.commit()
        bot.send_message(message.chat.id, f"✅ Do'kon saqlandi: {shop_name}")
    except: bot.send_message(message.chat.id, "❌ Bu do'kon allaqachon mavjud.")
    finally: conn.close()

@bot.message_handler(func=lambda message: message.text == "🛒 Yangi Buyurtma Urish")
def start_order(message):
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM shops")
    shops = cursor.fetchall()
    conn.close()
    if not shops:
        bot.send_message(message.chat.id, "❌ Do'konlar yo'q.")
        return
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    for s in shops: markup.add(types.KeyboardButton(s[0]))
    msg = bot.send_message(message.chat.id, "Do'konni tanlang:", reply_markup=markup)
    bot.register_next_step_handler(msg, choose_price_type)

def choose_price_type(message):
    user_steps[message.from_user.id] = {'shop_name': message.text, 'cart': {}, 'price_type': None}
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True).row(types.KeyboardButton("💰 Ulgurji (Optom)"), types.KeyboardButton("🛍 Chakana"))
    msg = bot.send_message(message.chat.id, "Narx turini tanlang:", reply_markup=markup)
    bot.register_next_step_handler(msg, show_products_to_agent)

def show_products_to_agent(message):
    p_type = "optom" if "Ulgurji" in message.text else "chakana"
    user_steps[message.from_user.id]['price_type'] = p_type
    send_product_list_menu(message)

def send_product_list_menu(message):
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM products")
    prods = cursor.fetchall()
    conn.close()
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    for p in prods: markup.add(types.KeyboardButton(p[0]))
    markup.add(types.KeyboardButton("✅ Buyurtmani yakunlash"))
    msg = bot.send_message(message.chat.id, "Mahsulotni tanlang:", reply_markup=markup)
    bot.register_next_step_handler(msg, ask_quantity)

def ask_quantity(message):
    if message.text == "✅ Buyurtmani yakunlash":
        finish_order(message)
        return
    user_steps[message.from_user.id]['current_product'] = message.text
    msg = bot.send_message(message.chat.id, f"🔢 Miqdorini kiriting:", reply_markup=types.ReplyKeyboardRemove())
    bot.register_next_step_handler(msg, add_to_cart)

def add_to_cart(message):
    uid = message.from_user.id
    try:
        qty = float(message.text)
        p_name = user_steps[uid]['current_product']
        user_steps[uid]['cart'][p_name] = qty
        bot.send_message(message.chat.id, f"📥 Qo'shildi: {p_name} - {qty}")
        send_product_list_menu(message)
    except:
        bot.send_message(message.chat.id, "❌ Faqat raqam kiriting.")
        send_product_list_menu(message)

def finish_order(message):
    uid = message.from_user.id
    data = user_steps.get(uid)
    if not data or not data['cart']:
        bot.send_message(message.chat.id, "Savat bo'sh!", reply_markup=get_main_menu("agent"))
        return
        
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    total_sum, items_text, p_type = 0, "", data['price_type']
    excel_cart_items = []
    
    for p_name, qty in data['cart'].items():
        cursor.execute("SELECT optom_price, chakana_price, stock, id FROM products WHERE name = ?", (p_name,))
        prod = cursor.fetchone()
        optom_p = prod[0] if prod[0] is not None else 0
        chakana_p = prod[1] if prod[1] is not None else 0
        stock_p = prod[2] if prod[2] is not None else 0
        
        price = optom_p if p_type == "optom" else chakana_p
        summa = price * qty
        total_sum += summa
        items_text += f"{p_name} - {qty}x = {summa:,.0f} so'm\n"
        cursor.execute("UPDATE products SET stock = ? WHERE id = ?", (stock_p - qty, prod[3]))
        excel_cart_items.append({'name': p_name, 'qty': qty, 'price': price})
        
    cursor.execute("SELECT debt FROM shops WHERE name = ?", (data['shop_name'],))
    shop_res = cursor.fetchone()
    current_debt = shop_res[0] if shop_res and shop_res[0] is not None else 0
    
    cursor.execute("UPDATE shops SET debt = ? WHERE name = ?", (current_debt + total_sum, data['shop_name']))
    cursor.execute("SELECT name FROM users WHERE tg_id = ?", (uid,))
    agent_res = cursor.fetchone()
    agent_name = agent_res[0] if agent_res else "Nomalum"
    bugun = datetime.now().strftime("%Y-%m-%d %H:%M")
    
    cursor.execute("INSERT INTO orders (shop_name, agent_name, total_sum, items_text, status, date) VALUES (?, ?, ?, ?, 'Yangi', ?)", (data['shop_name'], agent_name, total_sum, items_text, bugun))
    order_id = cursor.lastrowid
    conn.commit()
    conn.close()
    
    excel_file = create_excel_invoice(order_id, data['shop_name'], agent_name, bugun, p_type, excel_cart_items)
    
    bot.send_message(message.chat.id, f"✅ Buyurtma qabul qilindi! Jami: {total_sum:,.0f} so'm", reply_markup=get_main_menu("agent"))
    bot.send_message(ADMIN_ID, f"🔔 YANGI BUYURTMA:\n\nDo'kon: {data['shop_name']}\nSumma: {total_sum:,.0f} so'm")
    
    with open(excel_file, 'rb') as doc:
        bot.send_document(ADMIN_ID, doc, caption=f"📄 Nakladnoy (#{order_id})")
    if os.path.exists(excel_file): os.remove(excel_file)
    
    group_text = f"📝 <b>YANGI BUYURTMA</b>\n🏪 Do'kon: {data['shop_name']}\n\n{items_text}"
    try: bot.send_message(SEX_GROUP_ID, group_text, parse_mode="HTML")
    except: pass
    
    if uid in user_steps: del user_steps[uid]

@bot.callback_query_handler(func=lambda call: call.data.startswith("approve_"))
def approve_agent_cb(call):
    agent_id = int(call.data.split("_")[-1])
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("UPDATE users SET role = 'agent' WHERE tg_id = ?", (agent_id,))
    conn.commit()
    conn.close()
    bot.answer_callback_query(call.id, "Agent tasdiqlandi!")
    try: bot.send_message(agent_id, "🎉 Sizning so'rovingiz tasdiqlandi! /start bosing", reply_markup=get_main_menu("agent"))
    except: pass

@bot.message_handler(func=lambda message: message.text == "📝 Ro'yxatdan o'tish")
def register_start(message):
    msg = bot.send_message(message.chat.id, "Ismingizni kiriting:")
    bot.register_next_step_handler(msg, process_register_name)

def process_register_name(message):
    name = message.text
    msg = bot.send_message(message.chat.id, "Telefon raqamingizni kiriting:")
    bot.register_next_step_handler(msg, lambda m: save_pending_user(m, name))

def save_pending_user(message, name):
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO users (tg_id, name, phone, role) VALUES (?, ?, ?, 'pending')", (message.from_user.id, name, message.text))
    conn.commit()
    conn.close()
    bot.send_message(message.chat.id, "⏳ Admin tasdig'i kutilmoqda.")
    bot.send_message(ADMIN_ID, f"🔔 Yangi ro'yxatdan o'tgan: {name} ({message.text})")

@bot.message_handler(func=lambda message: message.text == "📜 Mening Buyurtmalarim")
def my_orders(message):
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM users WHERE tg_id = ?", (message.from_user.id,))
    res = cursor.fetchone()
    agent_name = res[0] if res else ""
    cursor.execute("SELECT shop_name, total_sum, status, date FROM orders WHERE agent_name = ? ORDER BY id DESC LIMIT 5", (agent_name,))
    orders = cursor.fetchall()
    conn.close()
    text = "<b>📜 Oxirgi buyurtmalar:</b>\n\n"
    for o in orders: 
        t_sum = o[1] if o[1] is not None else 0
        text += f"🏪 {o[0]} | {t_sum:,.0f} so'm | {o[2]} | {o[3]}\n\n"
    bot.send_message(message.chat.id, text, parse_mode="HTML")

@bot.message_handler(func=lambda message: message.text == "💰 Qarz/To'lov yozish")
def pay_debt_start(message):
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM shops")
    shops = cursor.fetchall()
    conn.close()
    if not shops:
        bot.send_message(message.chat.id, "❌ Do'konlar yo'q.")
        return
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    for s in shops: markup.add(types.KeyboardButton(s[0]))
    msg = bot.send_message(message.chat.id, "Do'konni tanlang:", reply_markup=markup)
    bot.register_next_step_handler(msg, ask_payment_amount)

def ask_payment_amount(message):
    shop_name = message.text
    msg = bot.send_message(message.chat.id, f"'{shop_name}' qancha to'lov qildi (summa):", reply_markup=types.ReplyKeyboardRemove())
    bot.register_next_step_handler(msg, process_payment, shop_name)

def process_payment(message, shop_name):
    try:
        amount = float(message.text)
        today = datetime.now().strftime("%Y-%m-%d %H:%M")
        conn = sqlite3.connect("xoji_aka_factory.db")
        cursor = conn.cursor()
        
        cursor.execute("SELECT debt FROM shops WHERE name = ?", (shop_name,))
        s_res = cursor.fetchone()
        curr_debt = s_res[0] if s_res and s_res[0] is not None else 0
        
        cursor.execute("UPDATE shops SET debt = ? WHERE name = ?", (curr_debt - amount, shop_name))
        cursor.execute("INSERT INTO incomes (source, amount, date) VALUES (?, ?, ?)", (f"Qarz to'lovi ({shop_name})", amount, today))
        conn.commit()
        conn.close()
        bot.send_message(message.chat.id, f"✅ To'lov yozildi: {amount:,.0f} so'm chegirildi.", reply_markup=get_main_menu("agent"))
    except: bot.send_message(message.chat.id, "❌ Faqat raqam kiriting.", reply_markup=get_main_menu("agent"))

# ================= FLASK WEB INTERFACE =================
@app.route('/')
def web_index():
    filter_date = request.args.get('filter_date')
    
    conn = sqlite3.connect("xoji_aka_factory.db")
    cursor = conn.cursor()
    cursor.execute("SELECT id, name, optom_price, chakana_price, stock FROM products")
    products = cursor.fetchall()
    
    cursor.execute("SELECT id, name, phone, debt FROM shops")
    shops = cursor.fetchall()
    
    if filter_date:
        cursor.execute("SELECT id, shop_name, agent_name, total_sum, status, date FROM orders WHERE date LIKE ? ORDER BY id DESC", (f"{filter_date}%",))
    else:
        cursor.execute("SELECT id, shop_name, agent_name, total_sum, status, date FROM orders ORDER BY id DESC")
    orders = cursor.fetchall()
    
    cursor.execute("SELECT SUM(debt) FROM shops")
    total_debt = cursor.fetchone()[0] or 0
    
    cursor.execute("SELECT SUM(amount) FROM incomes")
    total_income = cursor.fetchone()[0] or 0
    
    cursor.execute("SELECT SUM(amount) FROM expenses")
    total_expense = cursor.fetchone()[0] or 0
    
    kassa_balance = total_income - total_expense
    conn.close()
    
    WEB_HTML = """
    <!DOCTYPE html>
    <html lang="uz">
    <head>
        <meta charset="UTF-8">
        <title>Xoji Aka ERP — Web Panel</title>
        <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    </head>
    <body class="bg-light">
        <div class="container py-4">
            <h2 class="mb-4 text-primary fw-bold">🏭 Xoji Aka ERP Boshqaruv Paneli</h2>
            <div class="row g-3 mb-4">
                <div class="col-md-4">
                    <div class="p-3 bg-white shadow-sm rounded border">
                        <small class="text-muted">Kassadagi Naqd Pul</small>
                        <h4 class="fw-bold text-success">{{ "{:,.0f}".format(kassa_balance) }} so'm</h4>
                    </div>
                </div>
                <div class="col-md-4">
                    <div class="p-3 bg-white shadow-sm rounded border">
                        <small class="text-muted">Umumiy Nasiya (Qarz)</small>
                        <h4 class="fw-bold text-danger">{{ "{:,.0f}".format(total_debt) }} so'm</h4>
                    </div>
                </div>
            </div>
            
            <div class="bg-white p-4 shadow-sm rounded border mb-4">
                <h5 class="fw-bold mb-3">Buyurtmalar ro'yxati</h5>
                <table class="table table-striped">
                    <thead>
                        <tr><th>ID</th><th>Do'kon</th><th>Agent</th><th>Summa</th><th>Status</th><th>Sana</th></tr>
                    </thead>
                    <tbody>
                        {% for o in orders %}
                        <tr>
                            <td>#{{ o[0] }}</td>
                            <td>{{ o[1] }}</td>
                            <td>{{ o[2] }}</td>
                            <td>{{ "{:,.0f}".format(o[3]) }} so'm</td>
                            <td><span class="badge bg-secondary">{{ o[4] }}</span></td>
                            <td>{{ o[5] }}</td>
                        </tr>
                        {% else %}
                        <tr><td colspan="6" class="text-center text-muted">Buyurtmalar yo'q</td></tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
        </div>
    </body>
    </html>
    """
    return render_template_string(WEB_HTML, orders=orders, products=products, shops=shops, total_debt=total_debt, kassa_balance=kassa_balance)

def run_bot():
    bot.infinity_polling(skip_pending=True)

if __name__ == '__main__':
    init_db()
    auto_insert_products()
    
    # Telegram botni fonda ishga tushirish
    bot_thread = threading.Thread(target=run_bot, daemon=True)
    bot_thread.start()
    
    # Render.com beradigan portga moslab ishga tushirish
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
