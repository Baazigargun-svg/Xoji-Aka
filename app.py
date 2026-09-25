from datetime import datetime
import io
import os
import sqlite3
import threading
import time
from flask import (
    Flask,
    redirect,
    render_template_string,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests
import telebot
from telebot import types


def keep_alive():
  while True:
    try:
      url = "https://xoji-aka.onrender.com/"
      requests.get(url)
      print("Ping yuborildi, server uyg'oq!")
    except Exception as e:
      print(f"Ping xatosi: {e}")
    time.sleep(540)


threading.Thread(target=keep_alive, daemon=True).start()

# --- SOZLAMALAR ---
BOT_TOKEN = "8573337094:AAHR8_0QBjohJ3RXk7m2WtoqjuivLclVBNQ"
ADMIN_ID = 6851851908
SEX_GROUP_ID = -1003936599812

bot = telebot.TeleBot(BOT_TOKEN)
app = Flask(__name__)
DB_NAME = "xoji_aka_factory.db"
app.secret_key = "xoji_aka_maxfiy_kalit_2026"

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")

if not os.path.exists(UPLOAD_FOLDER):
  os.makedirs(UPLOAD_FOLDER)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
user_steps = {}


def get_db_connection():
  conn = sqlite3.connect(DB_NAME)
  conn.row_factory = sqlite3.Row
  return conn


def init_web_db():
  conn = get_db_connection()
  cursor = conn.cursor()

  cursor.execute(
      """CREATE TABLE IF NOT EXISTS users (tg_id INTEGER PRIMARY KEY, name TEXT, phone TEXT, role TEXT DEFAULT 'pending')"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS expenses 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, reason TEXT, amount REAL, date TEXT)"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS incomes 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT, amount REAL, date TEXT)"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS product_incomes 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, product_name TEXT, qty REAL, cost_price REAL, date TEXT)"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS orders 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, shop_name TEXT, agent_name TEXT, items_text TEXT, total_sum REAL, discount REAL DEFAULT 0, status TEXT, date TEXT, price_type TEXT, comment TEXT DEFAULT '')"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS order_status_history 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, order_id INTEGER, status TEXT, changed_at TEXT)"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS products 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE, category TEXT DEFAULT 'Boshqa', stock REAL DEFAULT 0, cost_price REAL DEFAULT 0, optom_price REAL DEFAULT 0, chakana_price REAL DEFAULT 0)"""
  )
  cursor.execute(
      """CREATE TABLE IF NOT EXISTS shops 
                      (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE, phone TEXT, debt REAL DEFAULT 0, visit_days TEXT, region TEXT DEFAULT '', landmark TEXT DEFAULT '', inventory TEXT DEFAULT '')"""
  )

  migrations = [
      (
          "ALTER TABLE products ADD COLUMN category TEXT DEFAULT 'Boshqa'",
          "category",
      ),
      ("ALTER TABLE products ADD COLUMN cost_price REAL DEFAULT 0", "cost_price"),
      (
          "ALTER TABLE products ADD COLUMN optom_price REAL DEFAULT 0",
          "optom_price",
      ),
      (
          "ALTER TABLE products ADD COLUMN chakana_price REAL DEFAULT 0",
          "chakana_price",
      ),
      ("ALTER TABLE products ADD COLUMN stock REAL DEFAULT 0", "stock"),
      ("ALTER TABLE shops ADD COLUMN visit_days TEXT DEFAULT ''", "visit_days"),
      ("ALTER TABLE shops ADD COLUMN region TEXT DEFAULT ''", "region"),
      ("ALTER TABLE shops ADD COLUMN landmark TEXT DEFAULT ''", "landmark"),
      ("ALTER TABLE shops ADD COLUMN inventory TEXT DEFAULT ''", "inventory"),
      ("ALTER TABLE orders ADD COLUMN discount REAL DEFAULT 0", "discount"),
      ("ALTER TABLE orders ADD COLUMN price_type TEXT", "price_type"),
      ("ALTER TABLE orders ADD COLUMN comment TEXT DEFAULT ''", "comment"),
  ]

  for query, col in migrations:
    try:
      cursor.execute(query)
    except:
      pass

  cursor.execute(
      "INSERT OR REPLACE INTO users (tg_id, name, phone, role) VALUES (?,"
      " 'Admin', '', 'admin')",
      (ADMIN_ID,),
  )

  initial_agents = [
      (8241020136, "Qoraboyev Sirojiddin", "+998935075540", "agent,sex"),
      (2101923750, "Qoraboyeva Charos", "+998940300206", "agent"),
  ]
  for ag_id, ag_name, ag_phone, ag_role in initial_agents:
    cursor.execute(
        "INSERT OR IGNORE INTO users (tg_id, name, phone, role) VALUES (?, ?,"
        " ?, ?)",
        (ag_id, ag_name, ag_phone, ag_role),
    )

  initial_shops = [
      ("Vanselling", "", "SEX Gulim"),
      ("Sherzod market", "", "Uzgazoil qatori"),
      ("Qayumov Kamol", "", "Murch boboga yetmasdan"),
      ("Akbar aka", "", "Bekat murch bobo"),
      ("Pub house", "", "Pub house"),
      ("Chapayev roparasi", "", "Chapayev roparasi"),
      ("Toshpoʻlat aka", "", "Uchrashuv yoni"),
      ("Abdulfayz bekat", "", ""),
      ("Dilorom", "", "Tutzor"),
      ("Joʻrabek aka", "", "Anorcha tagi"),
      ("Vohas", "", "Vohas"),
      ("Oʻzbegim market", "", "Davr bank yoni"),
      ("Pokiza market", "", "Muz saroy yoni"),
      ("Soxibkor doʻstlik", "", ""),
      ("Feruza non sex", "", "Nigoh non sexi yoni"),
      ("South brothers", "", "Doktor A qatori"),
      ("Gulmira opa", "", "Cola orqasi"),
      ("Taniqulov Oʻktam", "", ""),
      ("Abdulloh market", "", "Feredun café"),
      ("23-market", "", "23-sartarosh yoni"),
      ("Cola market", "", "South brothersga yetmay"),
      ("Muxlis Market", "", "Movaro"),
      ("Anjir Market", "", "Movaro"),
      ("Sevimli Market", "", "Yashil dunyo"),
      ("Abbos market", "", "Yashil dunyo"),
      ("Dilya opa", "", "Yashil dunyo"),
      ("Sanjar aka", "", "Yashil dunyo"),
      ("Nur market", "", "Yashil dunyo"),
      ("547 market", "", "Yashil dunyo"),
      ("Shox market", "", "Yashil dunyo"),
      ("Makro market", "", "Yashil dunyo"),
      ("Xusan bobo market", "", "Yashil dunyo"),
      ("Fresh market umid aka", "", "Movaro"),
      ("Pul hokim", "", "Adliya yoʻli"),
      ("Hoji ona market", "", "Med yoni"),
      ("Chinor market", "", "Med yoni"),
      ("16 market", "", "Begoyim roʻparasi"),
      ("Lada yoni", "", "Lada yoni"),
      ("Alibek aka", "", "Sohil pastlik"),
      ("4 aka-uka", "", "sohil"),
      ("Rayxon opa sohil", "", "Guliston ma-si"),
      ("Muhabbat opa", "", "Boyqishloq"),
      ("Moyka yoni", "", "Moyka yoni"),
      ("Malika yoni optom", "", "Malika yoni optom"),
      ("AR market", "", "Abdurashid market"),
      ("Kam-kam Market", "", "23-dom yoni"),
      ("Osiyo tagi", "", ""),
      ("Nigora opa", "", "Eski pioner oldi"),
      ("Sherbek aka/fresh M", "", "Best roʻparasi"),
      ("Otabek aka", "", "antena tagi"),
      ("Universal Market", "", "Lola kafe yoni"),
      ("Aka-Uka Market", "", "Masjid yoni"),
      ("Oila Market", "", "zilyonni yoʻli"),
      ("Boxo market", "", "zilyonni yoʻli"),
      ("Otabek zapchast M", "", "zilyonni yoʻli"),
      ("Maya market", "", "Gostsatndart yoni"),
      ("Umida opa", "", "Mashhura yoʻli"),
      ("Darband city", "", "Vokzal yoni"),
      ("Otajon market", "", "Tisudan keyin"),
      ("Barakali market", "", "Tisu yoni"),
      ("Asilabonu", "", "Boysun bekati yoni"),
      ("Norqulova Nargiza", "", "Hayit ala uyi taraf"),
      ("Sherzod aka", "", "Harbiy doʻkon"),
      ("Bahor Market", "", "Termiz tuman"),
      ("Baraka Market", "", "Termiz tuman"),
      ("Farxod Market", "", "Senter Kamaz"),
      ("Ariqcha Market", "", "Senter Kamaz"),
      ("Xolida Market", "", "Limon"),
  ]
  for s_name, s_phone, s_region in initial_shops:
    cursor.execute(
        """INSERT OR IGNORE INTO shops (name, phone, debt, region) VALUES (?, ?, 0, ?)""",
        (s_name, s_phone, s_region),
    )

  initial_products = [
      (
          "Pelmen /300 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          6200,
          14000.0,
          14000.0,
      ),
      (
          "Pelmen /500 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          9900,
          24000.0,
          24000.0,
      ),
      (
          "Pelmen rasepnoy /kg",
          "Yarim Tayyor Mahsulotlari",
          1000,
          19500,
          46000.0,
          46000.0,
      ),
      (
          "Teftel /300 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          12000,
          23000.0,
          23000.0,
      ),
      (
          "Pelmen ossarti /500 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          18000,
          30000.0,
          30000.0,
      ),
      (
          "Pelmen ossarti /300 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          11500,
          20000.0,
          20000.0,
      ),
      (
          "Golubtsi /500 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          14000,
          25000.0,
          25000.0,
      ),
      (
          "Tok doʻlma /300 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          11000.0,
          25000.0,
          25000.0,
      ),
      (
          "Karam doʻlma /300 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          11000.0,
          25000.0,
          25000.0,
      ),
      (
          "Somsa kesilgan /800 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          6000.0,
          17000.0,
          17000.0,
      ),
      (
          "Oʻrama xamir",
          "Yarim Tayyor Mahsulotlari",
          1000,
          6000.0,
          17000.0,
          17000.0,
      ),
      ("KFC", "Yarim Tayyor Mahsulotlari", 1000, 27500, 30000.0, 30000.0),
      (
          "KFC Gulim",
          "Yarim Tayyor Mahsulotlari",
          1000.0,
          15000,
          30000.0,
          30000.0,
      ),
      (
          "Osh masalliq 500 gr",
          "Yarim Tayyor Mahsulotlari",
          1000,
          6000.0,
          13000.0,
          13000.0,
      ),
      (
          "Osh masalliq 1 kg",
          "Yarim Tayyor Mahsulotlari",
          1000,
          7000.0,
          15000.0,
          15000.0,
      ),
      ("Lagʻmon", "Yarim Tayyor Mahsulotlari", 10000, 2000.0, 6000.0, 6000.0),
      ("Kotlet", "Yarim Tayyor Mahsulotlari", 1000, 11000.0, 25000.0, 25000.0),
      ("Lavash hamiri", "Yarim Tayyor Mahsulotlari", 1000, 3800.0, 7000.0, 7000.0),
      ("Manti hamiri", "Yarim Tayyor Mahsulotlari", 1000, 6000.0, 13000.0, 13000.0),
      ("Mirinda 250gr/30шт", "Yaxna ichimliklari", 1000, 7950.0, 9000.0, 9000.0),
      ("Mirinda 330gr/24шт", "Yaxna ichimliklari", 1000, 8800.0, 9000.0, 9000.0),
  ]
  for (
      p_name,
      p_cat,
      p_stock,
      p_cost,
      p_optom,
      p_chakana,
  ) in initial_products:
    cursor.execute(
        """INSERT OR IGNORE INTO products (name, category, stock, cost_price, optom_price, chakana_price) 
           VALUES (?, ?, ?, ?, ?, ?)""",
        (p_name, p_cat, p_stock, p_cost, p_optom, p_chakana),
    )

  conn.commit()
  conn.close()


def create_excel_invoice(
    order_id, shop_name, agent_name, date_str, price_type, cart_items, comment=""
):
  wb = Workbook()
  ws = wb.active
  ws.title = f"Nakladnoy_{order_id}"
  ws.sheet_view.showGridLines = True

  title_font = Font(name="Arial", size=16, bold=True)
  header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
  bold_font = Font(name="Arial", size=11, bold=True)
  header_fill = PatternFill(
      start_color="1F497D", end_color="1F497D", fill_type="solid"
  )
  total_fill = PatternFill(
      start_color="DCE6F1", end_color="DCE6F1", fill_type="solid"
  )
  thin_border = Border(
      left=Side(style="thin", color="B0B0B0"),
      right=Side(style="thin", color="B0B0B0"),
      top=Side(style="thin", color="B0B0B0"),
      bottom=Side(style="thin", color="B0B0B0"),
  )

  ws.merge_cells("A1:E1")
  ws["A1"] = "XOJI AKA FACTORY — NAKLADNOY"
  ws["A1"].font = title_font
  ws["A1"].alignment = Alignment(horizontal="center")

  ws["A3"] = f"Buyurtma ID: #{order_id}"
  ws["A3"].font = bold_font
  ws["D3"] = f"Sana: {date_str}"
  ws["A4"] = f"Do'kon (Klient): {shop_name}"
  ws["D4"] = f"Narx turi: {price_type.upper()}"
  ws["A5"] = f"Agent: {agent_name}"
  if comment:
    ws["A6"] = f"Izoh (Kommentariya): {comment}"
    ws["A6"].font = bold_font

  start_row = 8 if comment else 7

  headers = ["№", "Mahsulot nomi", "Miqdori", "Narxi (so'm)", "Jami summa"]
  for col_num, header_title in enumerate(headers, 1):
    cell = ws.cell(row=start_row, column=col_num, value=header_title)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

  row_num = start_row + 1
  total_sum = 0
  for idx, item in enumerate(cart_items, 1):
    ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(
        horizontal="center"
    )
    ws.cell(row=row_num, column=2, value=item["name"]).alignment = Alignment(
        horizontal="left"
    )
    ws.cell(row=row_num, column=3, value=item["qty"]).alignment = Alignment(
        horizontal="right"
    )
    ws.cell(row=row_num, column=4, value=item["price"]).alignment = Alignment(
        horizontal="right"
    )
    summa = item["qty"] * item["price"]
    total_sum += summa
    ws.cell(row=row_num, column=5, value=summa).alignment = Alignment(
        horizontal="right"
    )
    ws.cell(row=row_num, column=4).number_format = "#,##0"
    ws.cell(row=row_num, column=5).number_format = "#,##0"
    for col in range(1, 6):
      ws.cell(row=row_num, column=col).border = thin_border
    row_num += 1

  ws.merge_cells(
      start_row=row_num, start_column=1, end_row=row_num, end_column=4
  )
  ws.cell(row=row_num, column=1, value="JAMI TO'LOV:").alignment = Alignment(
      horizontal="right"
  )
  ws.cell(row=row_num, column=1).font = bold_font
  total_val = ws.cell(row=row_num, column=5, value=total_sum)
  total_val.font = bold_font
  total_val.number_format = "#,##0"
  for col in range(1, 6):
    ws.cell(row=row_num, column=col).fill = total_fill
    ws.cell(row=row_num, column=col).border = thin_border

  row_num += 2
  ws.cell(row=row_num, column=2, value="Qabul qildim: ____")
  ws.cell(row=row_num, column=4, value="Topshirdim: ____")

  file_name = f"Nakladnoy_{order_id}.xlsx"
  wb.save(file_name)
  return file_name


def create_excel_sex_income(income_id, staff_name, date_str, cart_items):
  wb = Workbook()
  ws = wb.active
  ws.title = f"Sex_Kirim_{income_id}"
  ws.sheet_view.showGridLines = True

  title_font = Font(name="Arial", size=16, bold=True)
  header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
  bold_font = Font(name="Arial", size=11, bold=True)
  header_fill = PatternFill(
      start_color="006100", end_color="006100", fill_type="solid"
  )
  total_fill = PatternFill(
      start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
  )
  thin_border = Border(
      left=Side(style="thin", color="B0B0B0"),
      right=Side(style="thin", color="B0B0B0"),
      top=Side(style="thin", color="B0B0B0"),
      bottom=Side(style="thin", color="B0B0B0"),
  )

  ws.merge_cells("A1:D1")
  ws["A1"] = "🏭 SEXGA MAHSULOT KIRIM (SKLAD)"
  ws["A1"].font = title_font
  ws["A1"].alignment = Alignment(horizontal="center")

  ws["A3"] = f"Kirim ID: #{income_id}"
  ws["A3"].font = bold_font
  ws["C3"] = f"Sana: {date_str}"
  ws["A4"] = f"Mas’ul xodim: {staff_name}"

  headers = ["№", "Mahsulot nomi", "Miqdori (Soni/Kg)"]
  for col_num, header_title in enumerate(headers, 1):
    cell = ws.cell(row=6, column=col_num, value=header_title)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

  row_num = 7
  total_qty = 0
  for idx, item in enumerate(cart_items, 1):
    ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(
        horizontal="center"
    )
    ws.cell(row=row_num, column=2, value=item["name"]).alignment = Alignment(
        horizontal="left"
    )
    ws.cell(row=row_num, column=3, value=item["qty"]).alignment = Alignment(
        horizontal="right"
    )
    total_qty += item["qty"]
    ws.cell(row=row_num, column=3).number_format = "#,##0.##"
    for col in range(1, 4):
      ws.cell(row=row_num, column=col).border = thin_border
    row_num += 1

  ws.merge_cells(
      start_row=row_num, start_column=1, end_row=row_num, end_column=2
  )
  ws.cell(row=row_num, column=1, value="JAMI QO'SHILGAN MIQDOR:").alignment = (
      Alignment(horizontal="right")
  )
  ws.cell(row=row_num, column=1).font = bold_font
  total_val = ws.cell(row=row_num, column=3, value=total_qty)
  total_val.font = bold_font
  total_val.number_format = "#,##0.##"
  for col in range(1, 4):
    ws.cell(row=row_num, column=col).fill = total_fill
    ws.cell(row=row_num, column=col).border = thin_border

  file_name = f"Sex_Kirim_{income_id}.xlsx"
  wb.save(file_name)
  return file_name


def get_main_menu(role):
  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  if role == "admin":
    markup.row(
        types.KeyboardButton("📦 Sklad & Mahsulotlar"),
        types.KeyboardButton("📊 Kunlik Hisobot"),
    )
    markup.row(
        types.KeyboardButton("👥 Agentlar boshqaruvi"),
        types.KeyboardButton("🏪 AKB (Do'konlar & Qarz)"),
    )
  elif role == "agent":
    markup.row(
        types.KeyboardButton("🛒 Yangi Buyurtma Urish"),
        types.KeyboardButton("🏪 Do'kon qo'shish (AKB)"),
    )
    markup.row(
        types.KeyboardButton("💰 Qarz/To'lov yozish"),
        types.KeyboardButton("📜 Mening Buyurtmalarim"),
    )
    markup.row(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))
  elif role == "sex":
    markup.row(
        types.KeyboardButton("📦 Skladga Mahsulot Kirim Qilish"),
        types.KeyboardButton("📋 Ombordagi Qoldiqlar"),
    )
    markup.row(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))
  else:
    markup.add(types.KeyboardButton("📝 Ro'yxatdan o'tish"))
  return markup


@bot.message_handler(commands=["start"])
def start_command(message):
  tg_id = message.from_user.id
  if tg_id == ADMIN_ID:
    bot.send_message(
        message.chat.id,
        "Xoji aka, xush kelibsiz! Boshqaruv paneli tayyor.",
        reply_markup=get_main_menu("admin"),
    )
    return
  conn = get_db_connection()
  user = conn.execute(
      "SELECT role, name FROM users WHERE tg_id = ?", (tg_id,)
  ).fetchone()
  conn.close()
  if user:
    role, name = user["role"], user["name"]
    if role == "pending":
      bot.send_message(
          message.chat.id,
          f"Salom {name}. So'rovingiz admin tasdig'ini kutyapti.",
      )
    elif "," in role:
      markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
      if "agent" in role:
        markup.add(types.KeyboardButton("👤 Agent rejimi"))
      if "sex" in role:
        markup.add(types.KeyboardButton("🏭 Sex rejimi"))
      bot.send_message(
          message.chat.id,
          f"Salom {name}! Iltimos, ish rejimini tanlang:",
          reply_markup=markup,
      )
    else:
      bot.send_message(
          message.chat.id,
          f"Salom {name}! Ishni boshlashimiz mumkin.",
          reply_markup=get_main_menu(role),
      )
  else:
    bot.send_message(
        message.chat.id,
        "Assalomu alaykum! Tizimga xush kelibsiz. Davom etish uchun"
        " ro'yxatdan o'ting.",
        reply_markup=get_main_menu("guest"),
    )


@bot.message_handler(
    func=lambda message: message.text
    in ["👤 Agent rejimi", "🏭 Sex rejimi", "🔄 Rolni almashtirish (Sex / Agent)"]
)
def switch_role_menu(message):
  tg_id = message.from_user.id
  conn = get_db_connection()
  user = conn.execute(
      "SELECT role, name FROM users WHERE tg_id = ?", (tg_id,)
  ).fetchone()
  conn.close()

  if not user:
    return

  role_str = user["role"]
  if message.text == "👤 Agent rejimi" or (
      "agent" in role_str
      and "sex" in role_str
      and message.text != "🏭 Sex rejimi"
  ):
    if message.text == "🔄 Rolni almashtirish (Sex / Agent)":
      markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
      markup.add(
          types.KeyboardButton("👤 Agent rejimi"),
          types.KeyboardButton("🏭 Sex rejimi"),
      )
      bot.send_message(
          message.chat.id, "Qaysi rejimga oʻtmoqchisiz?", reply_markup=markup
      )
      return

    bot.send_message(
        message.chat.id,
        "🛒 Agent rejimiga oʻtdingiz.",
        reply_markup=get_main_menu("agent"),
    )
  elif message.text == "🏭 Sex rejimi":
    bot.send_message(
        message.chat.id,
        "🏭 Sex rejimiga oʻtdingiz.",
        reply_markup=get_main_menu("sex"),
    )


@bot.message_handler(
    func=lambda message: message.text == "📦 Skladga Mahsulot Kirim Qilish"
)
def sex_income_start(message):
  conn = get_db_connection()
  prods = conn.execute("SELECT name FROM products").fetchall()
  conn.close()
  if not prods:
    bot.send_message(message.chat.id, "❌ Mahsulotlar topilmadi.")
    return

  user_steps[message.from_user.id] = {"sex_cart": {}}
  send_sex_product_menu(message)


def send_sex_product_menu(message):
  conn = get_db_connection()
  prods = conn.execute("SELECT name, stock FROM products").fetchall()
  conn.close()
  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  for p in prods:
    markup.add(types.KeyboardButton(p["name"]))
  markup.add(types.KeyboardButton("✅ Kirimni yakunlash"))
  markup.add(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))
  msg = bot.send_message(
      message.chat.id,
      "🏭 Kirim qilinadigan mahsulotni tanlang:",
      reply_markup=markup,
  )
  bot.register_next_step_handler(msg, sex_income_choose_product)


def sex_income_choose_product(message):
  if message.text == "🔄 Rolni almashtirish (Sex / Agent)":
    switch_role_menu(message)
    return
  if message.text == "✅ Kirimni yakunlash":
    finish_sex_income(message)
    return

  conn = get_db_connection()
  p_check = conn.execute(
      "SELECT id FROM products WHERE name = ?", (message.text,)
  ).fetchone()
  conn.close()

  if not p_check:
    bot.send_message(
        message.chat.id, "❌ Bunday mahsulot topilmadi. Ro'yxatdan tanlang:"
    )
    send_sex_product_menu(message)
    return

  user_steps[message.from_user.id]["current_product"] = message.text
  msg = bot.send_message(
      message.chat.id,
      f"🔢 <b>{message.text}</b> uchun miqdorni (kg / dona) kiriting:",
      parse_mode="HTML",
      reply_markup=types.ReplyKeyboardRemove(),
  )
  bot.register_next_step_handler(msg, sex_income_add_to_cart)


def sex_income_add_to_cart(message):
  uid = message.from_user.id
  try:
    qty = float(message.text)
    p_name = user_steps[uid]["current_product"]
    user_steps[uid]["sex_cart"][p_name] = qty
    bot.send_message(
        message.chat.id, f"📥 Qo'shildi: <b>{p_name}</b> — <b>{qty}</b>"
    )
    send_sex_product_menu(message)
  except:
    bot.send_message(message.chat.id, "❌ Xatolik! Faqat raqam kiriting.")
    send_sex_product_menu(message)


def finish_sex_income(message):
  uid = message.from_user.id
  data = user_steps.get(uid)
  if not data or not data.get("sex_cart"):
    bot.send_message(
        message.chat.id, "Savat bo'sh!", reply_markup=get_main_menu("sex")
    )
    return

  conn = get_db_connection()
  cursor = conn.cursor()
  items_text = ""
  cart_items = []

  for p_name, qty in data["sex_cart"].items():
    cursor.execute(
        "UPDATE products SET stock = stock + ? WHERE name = ?", (qty, p_name)
    )
    items_text += f"🔹 {p_name}: +{qty}\n"
    cart_items.append({"name": p_name, "qty": qty})

  staff_res = cursor.execute(
      "SELECT name FROM users WHERE tg_id = ?", (uid,)
  ).fetchone()
  staff_name = staff_res["name"] if staff_res else "Nomalum"
  bugun = datetime.now().strftime("%Y-%m-%d %H:%M")

  cursor.execute(
      "INSERT INTO product_incomes (product_name, qty, cost_price, date) VALUES"
      " (?, ?, 0, ?)",
      (items_text, sum(data["sex_cart"].values()), bugun),
  )
  income_id = cursor.lastrowid
  conn.commit()
  conn.close()

  excel_file = create_excel_sex_income(
      income_id, staff_name, bugun, cart_items
  )

  bot.send_message(
      message.chat.id,
      f"✅ <b>Skladga mahsulotlar muvaffaqiyatli kirim qilindi!</b>\n\n{items_text}",
      parse_mode="HTML",
      reply_markup=get_main_menu("sex"),
  )
  with open(excel_file, "rb") as doc:
    bot.send_document(
        message.chat.id, doc, caption=f"📄 Kirim Nakladnoy (#{income_id})"
    )

  try:
    group_text = (
        f"🏭 <b>YANGI SKLAD KIRIMI (#{income_id})</b>\n"
        f"👤 <b>Mas’ul:</b> {staff_name}\n"
        f"📅 <b>Sana:</b> {bugun}\n\n"
        f"<b>Kirim qilingan mahsulotlar:</b>\n{items_text}"
    )
    bot.send_message(SEX_GROUP_ID, group_text, parse_mode="HTML")
    with open(excel_file, "rb") as doc_group:
      bot.send_document(
          SEX_GROUP_ID, doc_group, caption=f"📄 Kirim Nakladnoy (#{income_id})"
      )
  except Exception as e:
    print("Sex guruhiga yuborish xatosi:", e)

  if os.path.exists(excel_file):
    os.remove(excel_file)
  if uid in user_steps:
    del user_steps[uid]


def send_auto_invoice(client_name, order_details_text):
  conn = get_db_connection()
  client = conn.execute(
      "SELECT telegram_id FROM clients WHERE name = ? OR shop_name = ?",
      (client_name, client_name),
  ).fetchone()
  conn.close()

  if client and client["telegram_id"]:
    tg_id = client["telegram_id"]
    try:
      bot.send_message(
          tg_id,
          f"📄 **Sizning buyurtmangiz chek-fakturasi:**\n\n{order_details_text}",
      )
      print(f"Chek muvaffaqiyatli jo‘natildi: {tg_id}")
    except Exception as e:
      print(f"Xabar yuborishda xatolik ({tg_id}): {e}")
  else:
    print(
        f"Diqqat: '{client_name}' bazadan topilmadi yoki Telegram ID si"
        " kiritilmagan!"
    )


@bot.message_handler(func=lambda message: message.text == "📋 Ombordagi Qoldiqlar")
def sex_view_stock(message):
  conn = get_db_connection()
  prods = conn.execute("SELECT name, stock FROM products").fetchall()
  conn.close()
  text = "📦 <b>Ombordagi qoldiqlar:</b>\n\n"
  for p in prods:
    stk = p["stock"] if p["stock"] is not None else 0
    text += f"🔹 {p['name']}: <b>{int(stk)}</b>\n"
  bot.send_message(
      message.chat.id, text, parse_mode="HTML", reply_markup=get_main_menu("sex")
  )


@bot.message_handler(func=lambda message: message.text == "📦 Sklad & Mahsulotlar")
def admin_sklad_menu(message):
  if message.from_user.id != ADMIN_ID:
    return
  conn = get_db_connection()
  prods = conn.execute("SELECT id, name, stock FROM products").fetchall()
  conn.close()

  inline_kb = types.InlineKeyboardMarkup(row_width=1)
  if not prods:
    bot.send_message(message.chat.id, "📦 Omborxonada mahsulotlar qolmagan.")
    return
  for p in prods:
    stock_val = p["stock"] if p["stock"] is not None else 0
    inline_kb.add(
        types.InlineKeyboardButton(
            f"🔹 {p['name']} ({int(stock_val)} kg/dona)",
            callback_data=f"adm_prod_{p['id']}",
        )
    )
  inline_kb.add(
      types.InlineKeyboardButton(
          "🆕 ✨ YANGI MAHSULOT QO'SHISH", callback_data="adm_create_product"
      )
  )
  bot.send_message(
      message.chat.id,
      "<b>📦 Sklad nazorati:</b>",
      parse_mode="HTML",
      reply_markup=inline_kb,
  )


@bot.callback_query_handler(func=lambda call: call.data.startswith("adm_prod_"))
def admin_product_detail(call):
  p_id = int(call.data.split("_")[-1])
  conn = get_db_connection()
  p = conn.execute(
      "SELECT name, optom_price, chakana_price, stock FROM products WHERE id ="
      " ?",
      (p_id,),
  ).fetchone()
  conn.close()
  if p:
    optom = p["optom_price"] if p["optom_price"] is not None else 0
    chakana = p["chakana_price"] if p["chakana_price"] is not None else 0
    stock = p["stock"] if p["stock"] is not None else 0
    text = (
        f"📦 <b>Mahsulot:</b> {p['name']}\n💰 Optom: {optom:,.0f} so'm\n🛍"
        f" Chakana: {chakana:,.0f} so'm\n🔢 <b>Qoldiq:</b> {int(stock)}"
    )
    markup = types.InlineKeyboardMarkup()
    markup.row(
        types.InlineKeyboardButton(
            "➕ Qoldiq Qo'shish", callback_data=f"stk_plus_{p_id}"
        ),
        types.InlineKeyboardButton(
            "➖ Qoldiq Ayirish", callback_data=f"stk_minus_{p_id}"
        ),
    )
    markup.row(
        types.InlineKeyboardButton(
            "💵 Optom Narx", callback_data=f"prc_optom_{p_id}"
        ),
        types.InlineKeyboardButton(
            "💵 Chakana Narx", callback_data=f"prc_chakana_{p_id}"
        ),
    )
    markup.row(
        types.InlineKeyboardButton("🗑 O'chirish", callback_data=f"stk_del_{p_id}")
    )
    bot.edit_message_text(
        text,
        call.message.chat.id,
        call.message.message_id,
        parse_mode="HTML",
        reply_markup=markup,
    )


@bot.callback_query_handler(
    func=lambda call: call.data.startswith(("stk_", "prc_"))
)
def admin_stock_price_actions(call):
  prefix, action, p_id = call.data.split("_")
  p_id = int(p_id)
  conn = get_db_connection()
  res = conn.execute("SELECT name FROM products WHERE id = ?", (p_id,)).fetchone()
  p_name = res["name"] if res else "Nomalum"
  conn.close()

  if action == "del":
    conn = get_db_connection()
    conn.execute("DELETE FROM products WHERE id = ?", (p_id,))
    conn.commit()
    conn.close()
    bot.send_message(
        call.message.chat.id, f"🗑 <b>{p_name}</b> o'chirildi.", parse_mode="HTML"
    )
    return

  msg = bot.send_message(
      call.message.chat.id,
      f"🔢 <b>{p_name}</b> uchun qiymat kiriting:",
      parse_mode="HTML",
  )
  bot.register_next_step_handler(msg, save_product_edits, prefix, action, p_id)


def save_product_edits(message, prefix, action, p_id):
  try:
    val = float(message.text)
    conn = get_db_connection()
    if action == "plus":
      conn.execute(
          "UPDATE products SET stock = stock + ? WHERE id = ?", (val, p_id)
      )
    elif action == "minus":
      conn.execute(
          "UPDATE products SET stock = stock - ? WHERE id = ?", (val, p_id)
      )
    elif action == "optom":
      conn.execute(
          "UPDATE products SET optom_price = ? WHERE id = ?", (val, p_id)
      )
    elif action == "chakana":
      conn.execute(
          "UPDATE products SET chakana_price = ? WHERE id = ?", (val, p_id)
      )
    conn.commit()
    conn.close()
    bot.send_message(message.chat.id, "✅ O'zgarish saqlandi!")
  except:
    bot.send_message(message.chat.id, "❌ Faqat raqam kiriting.")


@bot.callback_query_handler(func=lambda call: call.data == "adm_create_product")
def admin_create_product_start(call):
  msg = bot.send_message(
      call.message.chat.id, "📝 Yangi mahsulot NOMINI kiriting:"
  )
  bot.register_next_step_handler(msg, process_new_p_name)


def process_new_p_name(message):
  name = message.text
  msg = bot.send_message(
      message.chat.id, f"💰 '{name}' uchun Optom narxini kiriting:"
  )
  bot.register_next_step_handler(msg, process_new_p_optom, name)


def process_new_p_optom(message, name):
  try:
    optom = float(message.text)
    msg = bot.send_message(
        message.chat.id, f"🛍 '{name}' uchun Chakana narxini kiriting:"
    )
    bot.register_next_step_handler(msg, process_new_p_final, name, optom)
  except:
    bot.send_message(message.chat.id, "Xato! Faqat raqam kiriting.")


def process_new_p_final(message, name, optom):
  try:
    chakana = float(message.text)
    conn = get_db_connection()
    conn.execute(
        "INSERT INTO products (name, optom_price, chakana_price, stock)"
        " VALUES (?, ?, ?, 0)",
        (name, optom, chakana),
    )
    conn.commit()
    conn.close()
    bot.send_message(message.chat.id, f"✅ Yangi mahsulot qo'shildi: {name}")
  except:
    bot.send_message(message.chat.id, "Xato! Bu nom allaqachon mavjud.")


@bot.message_handler(
    func=lambda message: message.text
    in [
        "📊 Kunlik Hisobot",
        "🏪 AKB (Do'konlar & Qarz)",
        "👥 Agentlar boshqaruvi",
    ]
)
def admin_other_sections(message):
  if message.from_user.id != ADMIN_ID:
    return
  conn = get_db_connection()

  if message.text == "📊 Kunlik Hisobot":
    bugun = datetime.now().strftime("%Y-%m-%d")
    orders = conn.execute(
        "SELECT id, shop_name, total_sum, status FROM orders WHERE date LIKE ?",
        (f"{bugun}%",),
    ).fetchall()
    text = f"📊 <b>Bugungi buyurtmalar ({bugun}):</b>\n\n"
    inline_kb = types.InlineKeyboardMarkup()
    if not orders:
      text += "Hali buyurtma yo'q."
    for o in orders:
      t_sum = o["total_sum"] if o["total_sum"] is not None else 0
      text += (
          f"🆔 #{o['id']} | {o['shop_name']} | {t_sum:,.0f} so'm |"
          f" <b>{o['status']}</b>\n"
      )
      inline_kb.add(
          types.InlineKeyboardButton(
              f"⚙️ #{o['id']} Statusi", callback_data=f"mng_ord_{o['id']}"
          )
      )
    bot.send_message(
        message.chat.id, text, parse_mode="HTML", reply_markup=inline_kb
    )

  elif message.text == "🏪 AKB (Do'konlar & Qarz)":
    shops = conn.execute(
        "SELECT name, phone, debt, region FROM shops"
    ).fetchall()
    text = "<b>🏪 Do'konlar qarzlari:</b>\n\n"
    for s in shops:
      debt_val = s["debt"] if s["debt"] is not None else 0
      text += (
          f"🏢 {s['name']} ({s['region']}) — Qarz: {debt_val:,.0f} so'm\n"
      )
    bot.send_message(message.chat.id, text, parse_mode="HTML")

  elif message.text == "👥 Agentlar boshqaruvi":
    agents = conn.execute(
        "SELECT name, phone, role, tg_id FROM users WHERE role != 'admin'"
    ).fetchall()
    inline_kb = types.InlineKeyboardMarkup()
    text = "<b>👥 Agentlar va Sex xodimlari:</b>\n\n"
    for a in agents:
      status = (
          "✅ Faol"
          if any(r in a["role"] for r in ["agent", "sex"])
          else "⏳ Kutilmoqda"
      )
      text += f"👤 {a['name']} ({a['phone']}) - [{a['role']}] - {status}\n"
      if a["role"] == "pending":
        inline_kb.add(
            types.InlineKeyboardButton(
                f"👍 {a['name']}ni tasdiqlash",
                callback_data=f"approve_{a['tg_id']}",
            )
        )
    bot.send_message(
        message.chat.id, text, parse_mode="HTML", reply_markup=inline_kb
    )
  conn.close()


@bot.callback_query_handler(func=lambda call: call.data.startswith("mng_ord_"))
def manage_order_status_menu(call):
  order_id = int(call.data.split("_")[-1])
  markup = types.InlineKeyboardMarkup()
  markup.row(
      types.InlineKeyboardButton(
          "🔄 Yangi", callback_data=f"st_yangi_{order_id}"
      ),
      types.InlineKeyboardButton(
          "🚚 Otgruzka", callback_data=f"st_otgruzka_{order_id}"
      ),
  )
  markup.row(
      types.InlineKeyboardButton(
          "✅ Yetkazildi", callback_data=f"st_done_{order_id}"
      ),
      types.InlineKeyboardButton(
          "❌ Bekor", callback_data=f"st_otmen_{order_id}"
      ),
  )
  bot.send_message(
      call.message.chat.id,
      f"🆔 #{order_id} - Statusni tanlang:",
      reply_markup=markup,
  )


@bot.callback_query_handler(func=lambda call: call.data.startswith("st_"))
def change_status_logic(call):
  _, mode, order_id = call.data.split("_")
  order_id = int(order_id)
  status_map = {
      "yangi": "Yangi",
      "otgruzka": "Otgruzka",
      "done": "Yetkazildi",
      "otmen": "Bekor",
  }
  new_status = status_map[mode]
  conn = get_db_connection()
  conn.execute(
      "UPDATE orders SET status = ? WHERE id = ?", (new_status, order_id)
  )
  conn.commit()
  conn.close()
  bot.send_message(
      call.message.chat.id, f"✅ Status: <b>{new_status}</b>", parse_mode="HTML"
  )


@bot.callback_query_handler(func=lambda call: call.data.startswith("approve_"))
def approve_agent_cb(call):
  agent_id = int(call.data.split("_")[-1])
  conn = get_db_connection()
  conn.execute(
      "UPDATE users SET role = 'agent' WHERE tg_id = ?", (agent_id,)
  )
  conn.commit()
  conn.close()
  bot.answer_callback_query(call.id, "Foydalanuvchi tasdiqlandi!")
  try:
    bot.send_message(
        agent_id,
        "🎉 Sizning so'rovingiz tasdiqlandi! /start bosing",
        reply_markup=get_main_menu("agent"),
    )
  except:
    pass


@app.route("/add_shop", methods=["POST"])
def add_shop():
  name, phone, visit_days, region, landmark, inventory = (
      request.form["name"],
      request.form["phone"],
      request.form.get("visit_days", ""),
      request.form.get("region", ""),
      request.form.get("landmark", ""),
      request.form.get("inventory", ""),
  )
  conn = get_db_connection()
  try:
    conn.execute(
        "INSERT INTO shops (name, phone, debt, visit_days, region, landmark,"
        " inventory) VALUES (?, ?, 0, ?, ?, ?, ?)",
        (name, phone, visit_days, region, landmark, inventory),
    )
    conn.commit()
  except:
    pass
  conn.close()
  return redirect(url_for("operator_dashboard"))


@app.route("/update_shop/<int:shop_id>", methods=["POST"])
def update_shop(shop_id):
  name = request.form["name"]
  phone = request.form["phone"]
  region = request.form.get("region", "")
  landmark = request.form.get("landmark", "")
  visit_days = request.form.get("visit_days", "")
  inventory = request.form.get("inventory", "")

  conn = get_db_connection()
  try:
    conn.execute(
        "UPDATE shops SET name = ?, phone = ?, region = ?, landmark = ?,"
        " visit_days = ?, inventory = ? WHERE id = ?",
        (name, phone, region, landmark, visit_days, inventory, shop_id),
    )
    conn.commit()
  except Exception as e:
    print(e)
  conn.close()
  return redirect(url_for("operator_dashboard"))


@app.route("/update_product/<int:prod_id>", methods=["POST"])
def update_product(prod_id):
  name = request.form["name"]
  category = request.form.get("category", "Boshqa")
  stock = float(request.form.get("stock", 0))
  cost_price = float(request.form.get("cost_price", 0))
  optom_price = float(request.form.get("optom_price", 0))
  chakana_price = float(request.form.get("chakana_price", optom_price))

  conn = get_db_connection()
  try:
    conn.execute(
        "UPDATE products SET name = ?, category = ?, stock = ?, cost_price ="
        " ?, optom_price = ?, chakana_price = ? WHERE id = ?",
        (
            name,
            category,
            stock,
            cost_price,
            optom_price,
            chakana_price,
            prod_id,
        ),
    )
    conn.commit()
  except Exception as e:
    print("Mahsulotni tahrirlash xatosi:", e)
  conn.close()
  return redirect(url_for("operator_dashboard"))


@app.route("/adjust_stock_web", methods=["POST"])
def adjust_stock_web():
  prod_id = request.form.get("prod_id")
  action = request.form.get("action")
  qty = float(request.form.get("qty", 0))

  conn = get_db_connection()
  if action == "plus":
    conn.execute(
        "UPDATE products SET stock = stock + ? WHERE id = ?", (qty, prod_id)
    )
  elif action == "minus":
    conn.execute(
        "UPDATE products SET stock = stock - ? WHERE id = ?", (qty, prod_id)
    )
  conn.commit()
  conn.close()
  return redirect(url_for("operator_dashboard"))


@app.route("/add_agent_web", methods=["POST"])
def add_agent_web():
  name = request.form.get("name")
  phone = request.form.get("phone")
  tg_id = request.form.get("tg_id")
  role = request.form.get("role", "agent")
  if tg_id:
    try:
      tg_id = int(tg_id)
    except:
      tg_id = None

  if name and tg_id:
    conn = get_db_connection()
    try:
      conn.execute(
          "INSERT OR REPLACE INTO users (tg_id, name, phone, role) VALUES (?, ?,"
          " ?, ?)",
          (tg_id, name, phone or "", role),
      )
      conn.commit()
    except Exception as e:
      print("Agent qo'shish xatosi:", e)
    conn.close()
  return redirect(url_for("operator_dashboard"))


@app.route("/import_shops_excel", methods=["POST"])
def import_shops_excel():
  if "excel_file" not in request.files:
    return redirect(url_for("operator_dashboard"))
  file = request.files["excel_file"]
  if file.filename == "":
    return redirect(url_for("operator_dashboard"))

  try:
    df = pd.read_excel(file)
    conn = get_db_connection()
    cursor = conn.cursor()

    for _, row in df.iterrows():
      name = str(
          row.get("name", row.get("Do'kon nomi", row.get("Magazin", "")))
      ).strip()
      if not name or name == "nan":
        continue
      phone = str(row.get("phone", row.get("Telefon", row.get("Tel", "")))).strip()
      region = str(row.get("region", row.get("Hudud", row.get("Rayon", "")))).strip()
      landmark = str(
          row.get("landmark", row.get("Orienter", row.get("Mo'ljal", "")))
      ).strip()
      visit_days = str(
          row.get(
              "visit_days", row.get("Tashrif kunlari", row.get("Kunlar", ""))
          )
      ).strip()
      inventory = str(
          row.get(
              "inventory", row.get("Inventar", row.get("Jihoz", ""))
          )
      ).strip()

      debt_val = 0
      for d_key in ["debt", "Qarz", "Balans", "Borg"]:
        if d_key in row and pd.notna(row[d_key]):
          try:
            debt_val = float(row[d_key])
            break
          except:
            pass

      cursor.execute(
          """INSERT INTO shops (name, phone, debt, visit_days, region, landmark, inventory) 
                       VALUES (?, ?, ?, ?, ?, ?, ?)
                       ON CONFLICT(name) DO UPDATE SET 
                       phone=excluded.phone, region=excluded.region, landmark=excluded.landmark, 
                       visit_days=excluded.visit_days, inventory=excluded.inventory""",
          (
              name,
              phone if phone != "nan" else "",
              debt_val,
              visit_days if visit_days != "nan" else "",
              region if region != "nan" else "",
              landmark if landmark != "nan" else "",
              inventory if inventory != "nan" else "",
          ),
      )

    conn.commit()
    conn.close()
  except Exception as e:
    print("Excel import xatosi:", e)

  return redirect(url_for("operator_dashboard"))


@app.route("/import_products_excel", methods=["POST"])
def import_products_excel():
  if "excel_file" not in request.files:
    return redirect(url_for("operator_dashboard"))
  file = request.files["excel_file"]
  if file.filename == "":
    return redirect(url_for("operator_dashboard"))

  try:
    df = pd.read_excel(file)
    conn = get_db_connection()
    cursor = conn.cursor()

    for _, row in df.iterrows():
      name = str(
          row.get("name", row.get("Mahsulot", row.get("Tovar nomi", "")))
      ).strip()
      if not name or name == "nan":
        continue
      category = str(
          row.get("category", row.get("Kategoriya", "Boshqa"))
      ).strip()
      if not category or category == "nan":
        category = "Boshqa"

      stock_val = 0
      for s_key in ["stock", "Qoldiq", "Soni", "Miqdor"]:
        if s_key in row and pd.notna(row[s_key]):
          try:
            stock_val = float(row[s_key])
            break
          except:
            pass

      cost_val = 0
      for c_key in ["cost_price", "Tannarx", "Tan narx"]:
        if c_key in row and pd.notna(row[c_key]):
          try:
            cost_val = float(row[c_key])
            break
          except:
            pass

      optom_val = 0
      for o_key in ["optom_price", "Optom", "Optom narx"]:
        if o_key in row and pd.notna(row[o_key]):
          try:
            optom_val = float(row[o_key])
            break
          except:
            pass

      chakana_val = optom_val
      for ch_key in ["chakana_price", "Chakana", "Chakana narx"]:
        if ch_key in row and pd.notna(row[ch_key]):
          try:
            chakana_val = float(row[ch_key])
            break
          except:
            pass

      cursor.execute(
          """INSERT INTO products (name, category, stock, cost_price, optom_price, chakana_price) 
                       VALUES (?, ?, ?, ?, ?, ?)
                       ON CONFLICT(name) DO UPDATE SET 
                       category=excluded.category, stock=excluded.stock, 
                       cost_price=excluded.cost_price, optom_price=excluded.optom_price, 
                       chakana_price=excluded.chakana_price""",
          (name, category, stock_val, cost_val, optom_val, chakana_val),
      )

    conn.commit()
    conn.close()
  except Exception as e:
    print("Sklad Excel import xatosi:", e)

  return redirect(url_for("operator_dashboard"))


@app.route("/approve_agent/<int:tg_id>", methods=["POST"])
def web_approve_agent(tg_id):
  conn = get_db_connection()
  conn.execute("UPDATE users SET role = 'agent' WHERE tg_id = ?", (tg_id,))
  conn.commit()
  conn.close()
  try:
    bot.send_message(
        tg_id,
        "🎉 Sizning so'rovingiz tasdiqlandi! /start bosing",
        reply_markup=get_main_menu("agent"),
    )
  except:
    pass
  return redirect(url_for("operator_dashboard"))


@app.route("/delete_agent/<int:tg_id>", methods=["POST"])
def web_delete_agent(tg_id):
  conn = get_db_connection()
  conn.execute("DELETE FROM users WHERE tg_id = ?", (tg_id,))
  conn.commit()
  conn.close()
  return redirect(url_for("operator_dashboard"))


@bot.message_handler(func=lambda message: message.text == "🛒 Yangi Buyurtma Urish")
def start_order(message):
  conn = get_db_connection()
  shops = conn.execute("SELECT name FROM shops").fetchall()
  conn.close()
  if not shops:
    bot.send_message(message.chat.id, "❌ Do'konlar yo'q.")
    return
  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  for s in shops:
    markup.add(types.KeyboardButton(s["name"]))
  markup.add(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))
  msg = bot.send_message(message.chat.id, "Do'konni tanlang:", reply_markup=markup)
  bot.register_next_step_handler(msg, choose_price_type)


def choose_price_type(message):
  if message.text == "🔄 Rolni almashtirish (Sex / Agent)":
    switch_role_menu(message)
    return
  user_steps[message.from_user.id] = {
      "shop_name": message.text,
      "cart": {},
      "price_type": None,
  }
  markup = types.ReplyKeyboardMarkup(resize_keyboard=True).row(
      types.KeyboardButton("💰 Ulgurji (Optom)"),
      types.KeyboardButton("🛍 Chakana"),
  )
  markup.add(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))
  msg = bot.send_message(
      message.chat.id, "Narx turini tanlang:", reply_markup=markup
  )
  bot.register_next_step_handler(msg, show_categories_to_agent)


def show_categories_to_agent(message):
  if message.text == "🔄 Rolni almashtirish (Sex / Agent)":
    switch_role_menu(message)
    return
  p_type = "optom" if "Ulgurji" in message.text else "chakana"
  user_steps[message.from_user.id]["price_type"] = p_type

  conn = get_db_connection()
  categories = conn.execute(
      "SELECT DISTINCT category FROM products"
  ).fetchall()
  conn.close()

  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  for c in categories:
    cat_name = c["category"] if c["category"] else "Boshqa"
    markup.add(types.KeyboardButton(f"📁 {cat_name}"))
  markup.add(types.KeyboardButton("✅ Buyurtmani yakunlash"))
  markup.add(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))

  msg = bot.send_message(
      message.chat.id,
      "📁 Mahsulot kategoriyasini tanlang yoki buyurtmani yakunlang:",
      reply_markup=markup,
  )
  bot.register_next_step_handler(msg, choose_product_category)


def choose_product_category(message):
  if message.text == "🔄 Rolni almashtirish (Sex / Agent)":
    switch_role_menu(message)
    return
  if message.text == "✅ Buyurtmani yakunlash":
    ask_order_comment(message)
    return

  cat_name = message.text.replace("📁 ", "").strip()
  user_steps[message.from_user.id]["current_category"] = cat_name

  conn = get_db_connection()
  prods = conn.execute(
      "SELECT name FROM products WHERE category = ?", (cat_name,)
  ).fetchall()
  conn.close()

  if not prods:
    conn = get_db_connection()
    prods = conn.execute("SELECT name FROM products").fetchall()
    conn.close()

  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  for p in prods:
    markup.add(types.KeyboardButton(p["name"]))
  markup.add(types.KeyboardButton("🔙 Kategoriyalarga qaytish"))
  markup.add(types.KeyboardButton("✅ Buyurtmani yakunlash"))

  msg = bot.send_message(
      message.chat.id,
      f"📦 <b>{cat_name}</b> kategoriyasidagi mahsulotni tanlang:",
      parse_mode="HTML",
      reply_markup=markup,
  )
  bot.register_next_step_handler(msg, ask_quantity_or_navigation)


def ask_quantity_or_navigation(message):
  if message.text == "🔙 Kategoriyalarga qaytish":
    show_categories_to_agent_again(message)
    return
  if message.text == "✅ Buyurtmani yakunlash":
    ask_order_comment(message)
    return

  conn = get_db_connection()
  p_check = conn.execute(
      "SELECT id FROM products WHERE name = ?", (message.text,)
  ).fetchone()
  conn.close()

  if not p_check:
    bot.send_message(
        message.chat.id,
        "❌ Bunday mahsulot topilmadi. Iltimos, tugmalardan tanlang:",
    )
    choose_product_category(message)
    return

  user_steps[message.from_user.id]["current_product"] = message.text
  msg = bot.send_message(
      message.chat.id,
      f"🔢 <b>{message.text}</b> miqdorini kiriting:",
      parse_mode="HTML",
      reply_markup=types.ReplyKeyboardRemove(),
  )
  bot.register_next_step_handler(msg, add_to_cart)


def show_categories_to_agent_again(message):
  uid = message.from_user.id
  conn = get_db_connection()
  categories = conn.execute(
      "SELECT DISTINCT category FROM products"
  ).fetchall()
  conn.close()

  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  for c in categories:
    cat_name = c["category"] if c["category"] else "Boshqa"
    markup.add(types.KeyboardButton(f"📁 {cat_name}"))
  markup.add(types.KeyboardButton("✅ Buyurtmani yakunlash"))

  msg = bot.send_message(
      message.chat.id, "📁 Kategoriyani tanlang:", reply_markup=markup
  )
  bot.register_next_step_handler(msg, choose_product_category)


def add_to_cart(message):
  uid = message.from_user.id
  try:
    qty = float(message.text)
    p_name = user_steps[uid]["current_product"]
    user_steps[uid]["cart"][p_name] = qty
    bot.send_message(message.chat.id, f"📥 Qo'shildi: {p_name} - {qty}")

    cat_name = user_steps[uid].get("current_category", "Boshqa")
    conn = get_db_connection()
    prods = conn.execute(
        "SELECT name FROM products WHERE category = ?", (cat_name,)
    ).fetchall()
    conn.close()

    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    for p in prods:
      markup.add(types.KeyboardButton(p["name"]))
    markup.add(types.KeyboardButton("🔙 Kategoriyalarga qaytish"))
    markup.add(types.KeyboardButton("✅ Buyurtmani yakunlash"))

    msg = bot.send_message(
        message.chat.id,
        "Yana mahsulot qo'shasizmi yoki buyurtmani yakunlaysizmi?",
        reply_markup=markup,
    )
    bot.register_next_step_handler(msg, ask_quantity_or_navigation)
  except:
    bot.send_message(message.chat.id, "❌ Xatolik! Faqat raqam kiriting.")
    show_categories_to_agent_again(message)


def ask_order_comment(message):
  uid = message.from_user.id
  data = user_steps.get(uid)
  if not data or not data.get("cart"):
    bot.send_message(
        message.chat.id, "Savat bo'sh!", reply_markup=get_main_menu("agent")
    )
    return

  msg = bot.send_message(
      message.chat.id,
      "📝 Buyurtma uchun **kommentariya (izoh)** yozing (masalan: <i>'Ertalabga"
      " yetkazilsin'</i>):",
      parse_mode="HTML",
      reply_markup=types.ReplyKeyboardRemove(),
  )
  bot.register_next_step_handler(msg, finish_order_with_comment)


def finish_order_with_comment(message):
  uid = message.from_user.id
  comment = message.text if message.text else ""
  data = user_steps.get(uid)

  if not data or not data["cart"]:
    bot.send_message(
        message.chat.id, "Savat bo'sh!", reply_markup=get_main_menu("agent")
    )
    return

  conn = get_db_connection()
  total_sum, items_text, p_type = 0, "", data["price_type"]
  excel_cart_items = []

  for p_name, qty in data["cart"].items():
    prod = conn.execute(
        "SELECT optom_price, chakana_price, stock, id FROM products WHERE name ="
        " ?",
        (p_name,),
    ).fetchone()
    optom_p = prod["optom_price"] if prod["optom_price"] is not None else 0
    chakana_p = (
        prod["chakana_price"] if prod["chakana_price"] is not None else 0
    )
    stock_p = prod["stock"] if prod["stock"] is not None else 0

    price = optom_p if p_type == "optom" else chakana_p
    summa = price * qty
    total_sum += summa
    items_text += f"{p_name} - {qty}x = {summa:,.0f} so'm\n"
    conn.execute(
        "UPDATE products SET stock = ? WHERE id = ?",
        (stock_p - qty, prod["id"]),
    )
    excel_cart_items.append({"name": p_name, "qty": qty, "price": price})

  agent_res = conn.execute(
      "SELECT name FROM users WHERE tg_id = ?", (uid,)
  ).fetchone()
  agent_name = agent_res["name"] if agent_res else "Nomalum"
  bugun = datetime.now().strftime("%Y-%m-%d %H:%M")

  cursor = conn.cursor()
  cursor.execute(
      "INSERT INTO orders (shop_name, agent_name, total_sum, items_text,"
      " status, date, price_type, comment) VALUES (?, ?, ?, ?, 'Yangi', ?, ?,"
      " ?)",
      (
          data["shop_name"],
          agent_name,
          total_sum,
          items_text,
          bugun,
          p_type,
          comment,
      ),
  )
  order_id = cursor.lastrowid
  conn.commit()
  conn.close()

  excel_file = create_excel_invoice(
      order_id,
      data["shop_name"],
      agent_name,
      bugun,
      p_type,
      excel_cart_items,
      comment,
  )

  bot.send_message(
      message.chat.id,
      f"✅ Buyurtma qabul qilindi! Jami: {total_sum:,.0f} so'm",
      reply_markup=get_main_menu("agent"),
  )
  bot.send_message(
      ADMIN_ID,
      f"🔔 YANGI BUYURTMA (#{order_id}):\n\nDo'kon: {data['shop_name']}\nSumma:"
      f" {total_sum:,.0f} so'm\nIzoh: {comment}",
  )

  with open(excel_file, "rb") as doc:
    bot.send_document(ADMIN_ID, doc, caption=f"📄 Nakladnoy (#{order_id})")
  if os.path.exists(excel_file):
    os.remove(excel_file)

  try:
    group_text = (
        f"📝 <b>YANGI BUYURTMA (#{order_id})</b>\n"
        f"🏪 <b>Do'kon:</b> {data['shop_name']}\n"
        f"👤 <b>Agent:</b> {agent_name}\n"
        f"💬 <b>Izoh:</b> {comment}\n"
        f"💰 <b>Jami summa:</b> {total_sum:,.0f} so'm\n\n"
        f"<b>Mahsulotlar:</b>\n{items_text}"
    )
    bot.send_message(SEX_GROUP_ID, group_text, parse_mode="HTML")
    with open(
        create_excel_invoice(
            order_id,
            data["shop_name"],
            agent_name,
            bugun,
            p_type,
            excel_cart_items,
            comment,
        ),
        "rb",
    ) as doc_group:
      bot.send_document(
          SEX_GROUP_ID, doc_group, caption=f"📄 Nakladnoy (#{order_id})"
      )
  except Exception as e:
    print("Guruhga yuborish xatosi:", e)

  if uid in user_steps:
    del user_steps[uid]


@bot.message_handler(func=lambda message: message.text == "📝 Ro'yxatdan o'tish")
def register_start(message):
  msg = bot.send_message(message.chat.id, "Ismingizni kiriting:")
  bot.register_next_step_handler(msg, process_register_name)


def process_register_name(message):
  name = message.text
  msg = bot.send_message(message.chat.id, "Telefon raqamingizni kiriting:")
  bot.register_next_step_handler(msg, lambda m: save_pending_user(m, name))


def save_pending_user(message, name):
  conn = get_db_connection()
  conn.execute(
      "INSERT OR REPLACE INTO users (tg_id, name, phone, role) VALUES (?, ?, ?,"
      " 'pending')",
      (message.from_user.id, name, message.text),
  )
  conn.commit()
  conn.close()
  bot.send_message(message.chat.id, "⏳ Admin tasdig'i kutilmoqda.")
  bot.send_message(
      ADMIN_ID, f"🔔 Yangi ro'yxatdan o'tgan: {name} ({message.text})"
  )


@bot.message_handler(func=lambda message: message.text == "📜 Mening Buyurtmalarim")
def my_orders(message):
  conn = get_db_connection()
  res = conn.execute(
      "SELECT name FROM users WHERE tg_id = ?", (message.from_user.id,)
  ).fetchone()
  agent_name = res["name"] if res else ""
  orders = conn.execute(
      "SELECT shop_name, total_sum, status, date FROM orders WHERE agent_name ="
      " ? ORDER BY id DESC LIMIT 5",
      (agent_name,),
  ).fetchall()
  conn.close()
  text = "<b>📜 Oxirgi buyurtmalar:</b>\n\n"
  for o in orders:
    t_sum = o["total_sum"] if o["total_sum"] is not None else 0
    text += f"🏪 {o['shop_name']} | {t_sum:,.0f} so'm | {o['status']} | {o['date']}\n\n"
  bot.send_message(message.chat.id, text, parse_mode="HTML")


@bot.message_handler(func=lambda message: message.text == "💰 Qarz/To'lov yozish")
def pay_debt_start(message):
  conn = get_db_connection()
  shops = conn.execute("SELECT name FROM shops").fetchall()
  conn.close()
  if not shops:
    bot.send_message(message.chat.id, "❌ Do'konlar yo'q.")
    return
  markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
  for s in shops:
    markup.add(types.KeyboardButton(s["name"]))
  markup.add(types.KeyboardButton("🔄 Rolni almashtirish (Sex / Agent)"))
  msg = bot.send_message(message.chat.id, "Do'konni tanlang:", reply_markup=markup)
  bot.register_next_step_handler(msg, ask_payment_amount)


def ask_payment_amount(message):
  if message.text == "🔄 Rolni almashtirish (Sex / Agent)":
    switch_role_menu(message)
    return
  shop_name = message.text
  msg = bot.send_message(
      message.chat.id,
      f"'{shop_name}' qancha to'lov kirdi (summa):",
      reply_markup=types.ReplyKeyboardRemove(),
  )
  bot.register_next_step_handler(msg, process_payment, shop_name)


def process_payment(message, shop_name):
  try:
    amount = float(message.text)
    today = datetime.now().strftime("%Y-%m-%d %H:%M")
    conn = get_db_connection()
    s_res = conn.execute(
        "SELECT debt FROM shops WHERE name = ?", (shop_name,)
    ).fetchone()
    curr_debt = s_res["debt"] if s_res and s_res["debt"] is not None else 0

    conn.execute(
        "UPDATE shops SET debt = ? WHERE name = ?",
        (curr_debt - amount, shop_name),
    )
    conn.execute(
        "INSERT INTO incomes (source, amount, date) VALUES (?, ?, ?)",
        (f"Qarz to'lovi ({shop_name})", amount, today),
    )
    conn.commit()
    conn.close()
    bot.send_message(
        message.chat.id,
        f"✅ To'lov yozildi: {amount:,.0f} so'm chegirildi.",
        reply_markup=get_main_menu("agent"),
    )
  except:
    bot.send_message(
        message.chat.id,
        "❌ Faqat raqam kiriting.",
        reply_markup=get_main_menu("agent"),
    )


HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="uz">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Xoji Aka ERP — Premium Boshqaruv</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap-icons@1.11.0/font/bootstrap-icons.css">
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        :root { --bg-main: #f8fafc; --sidebar-bg: #0f172a; --primary-color: #3b82f6; }
        body { background-color: var(--bg-main); font-family: 'Inter', system-ui, sans-serif; color: #1e293b; }
        .sidebar { width: 150px; background: var(--sidebar-bg); min-height: 100vh; box-shadow: 4px 0 20px rgba(0,0,0,0.05); }
        .sidebar .nav-link { text-align: center; padding: 12px 6px; color: #94a3b8; font-size: 12px; font-weight: 500; border-radius: 10px; margin: 6px 8px; transition: all 0.25s ease; }
        .sidebar .nav-link:hover { background: rgba(255, 255, 255, 0.08); color: #ffffff; }
        .sidebar .nav-link.active { background: var(--primary-color); color: #ffffff; font-weight: 600; }
        .sidebar .nav-link i { font-size: 20px; display: block; margin-bottom: 4px; }
        .brand-logo-container { text-align: center; padding: 16px 5px 12px 5px; border-bottom: 1px solid rgba(255, 255, 255, 0.1); margin-bottom: 10px; }
        .brand-xa { font-family: 'Georgia', serif; font-weight: 900; font-size: 34px; line-height: 0.8; color: #ffffff; font-style: italic; }
        .brand-line { height: 3px; background-color: #ef4444; width: 60px; margin: 6px auto; border-radius: 2px; }
        .brand-name { font-weight: 800; font-size: 9px; letter-spacing: 2px; text-transform: uppercase; color: #cbd5e1; }
        .top-bar { background: #ffffff; border-bottom: 1px solid #e2e8f0; padding: 12px 25px; }
        .card-glass { background: #ffffff; border-radius: 16px; border: 1px solid #e2e8f0; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.02); }
        .stat-box { border-radius: 14px; padding: 22px; color: white; position: relative; overflow: hidden; }
        .stat-blue { background: linear-gradient(135deg, #3b82f6, #1d4ed8); }
        .stat-green { background: linear-gradient(135deg, #10b981, #047857); }
        .stat-red { background: linear-gradient(135deg, #ef4444, #b91c1c); }
        .table-custom th { font-weight: 600; font-size: 13px; background: #f8fafc; color: #475569; }
        .table-custom td { font-size: 13px; vertical-align: middle; }
        .day-badge { display: inline-block; padding: 5px 8px; margin: 2px; border-radius: 6px; font-size: 11px; font-weight: bold; background: #e2e8f0; color: #475569; cursor: pointer; border: 1px solid #cbd5e1; user-select: none; }
        .day-badge.selected { background: #3b82f6; color: white; border-color: #2563eb; }
        .shop-row { cursor: pointer; }
        .shop-row:hover { background-color: #f1f5f9 !important; }
        .product-row { cursor: pointer; }
        .product-row:hover { background-color: #f1f5f9 !important; }
    </style>
</head>
<body>
<div class="d-flex">
    <div class="sidebar d-flex flex-column flex-shrink-0 nav nav-pills" id="v-pills-tab" role="tablist">
        <div class="brand-logo-container">
            <div class="brand-xa">XA</div>
            <div class="brand-line"></div>
            <div class="brand-name">Xoji Aka</div>
        </div>
        <button class="nav-link active" data-bs-toggle="pill" data-bs-target="#tab-dashboard" type="button"><i class="bi bi-grid-fill"></i>Bosh Panel</button>
        <button class="nav-link" data-bs-toggle="pill" data-bs-target="#tab-orders" type="button"><i class="bi bi-cart-fill"></i>Buyurtmalar</button>
        <button class="nav-link" data-bs-toggle="pill" data-bs-target="#tab-inventory" type="button"><i class="bi bi-boxes"></i>Sklad</button>
        <button class="nav-link" data-bs-toggle="pill" data-bs-target="#tab-clients" type="button"><i class="bi bi-shop"></i>Do'konlar</button>
        <button class="nav-link" data-bs-toggle="pill" data-bs-target="#tab-agents" type="button"><i class="bi bi-people-fill"></i>Agentlar</button>
        <button class="nav-link" data-bs-toggle="pill" data-bs-target="#tab-kassa" type="button"><i class="bi bi-wallet2"></i>Kassa</button>
        <button class="nav-link" data-bs-toggle="pill" data-bs-target="#tab-reports" type="button"><i class="bi bi-file-earmark-bar-graph"></i>Hisobot</button>
    </div>
    <div class="flex-grow-1">
        <div class="top-bar d-flex justify-content-between align-items-center">
            <div class="fw-bold text-dark fs-6"><i class="bi bi-shield-check text-primary me-2"></i>Boshqaruv Markazi</div>
            <div class="d-flex align-items-center gap-2">
                <a href="/logout" class="btn btn-sm btn-outline-danger fw-bold"><i class="bi bi-box-arrow-right me-1"></i>Chiqish</a>
                <a href="/export_excel" class="btn btn-sm btn-success fw-bold"><i class="bi bi-file-earmark-excel me-1"></i>Excelga Yuklab Olish</a>
                <form method="GET" action="/" class="d-flex align-items-center gap-1 m-0 bg-light p-1 rounded border">
                    <span class="text-muted small px-1">Dan:</span>
                    <input type="date" name="start_date" value="{{ start_date }}" class="form-control form-control-sm" style="width: 130px;">
                    <span class="text-muted small px-1">Gacha:</span>
                    <input type="date" name="end_date" value="{{ end_date }}" class="form-control form-control-sm" style="width: 130px;">
                    <button type="submit" class="btn btn-sm btn-primary">Saralash</button>
                    <a href="/" class="btn btn-sm btn-light border" title="Tozalash">✕</a>
                </form>
            </div>
        </div>
        <div class="p-4">
            <div class="tab-content">
                <div class="tab-pane fade show active" id="tab-dashboard">
                    <div class="row g-4 mb-4">
                        <div class="col-md-4">
                            <div class="stat-box stat-blue shadow-sm">
                                <div class="small text-white-50">Tanlangan Davr Tushumi</div>
                                <h2 class="fw-bold mt-1 mb-0">{{ "{:,.0f}".format(daily_sum) }} <span class="fs-6">so'm</span></h2>
                            </div>
                        </div>
                        <div class="col-md-4">
                            <div class="stat-box stat-green shadow-sm">
                                <div class="small text-white-50">Kassadagi Naqd Pul</div>
                                <h2 class="fw-bold mt-1 mb-0">{{ "{:,.0f}".format(kassa_balance) }} <span class="fs-6">so'm</span></h2>
                            </div>
                        </div>
                        <div class="col-md-4">
                            <div class="stat-box stat-red shadow-sm">
                                <div class="small text-white-50">Umumiy Nasiya (Qarzlar)</div>
                                <h2 class="fw-bold mt-1 mb-0">{{ "{:,.0f}".format(total_debt) }} <span class="fs-6">so'm</span></h2>
                            </div>
                        </div>
                    </div>
                </div>
            </div>
        </div>
    </div>
</div>
</body>
</html>
"""

# Agar shablonlarni alohida funksiyada render qilish kerak bo'lsa:
@app.route("/")
def operator_dashboard():
  conn = get_db_connection()
  shops = conn.execute("SELECT * FROM shops").fetchall()
  products = conn.execute("SELECT * FROM products").fetchall()
  orders = conn.execute(
      "SELECT * FROM orders ORDER BY id DESC LIMIT 50"
  ).fetchall()
  conn.close()

  return render_template_string(
      HTML_TEMPLATE,
      shops=shops,
      products=products,
      orders=orders,
      daily_sum=0,
      kassa_balance=0,
      total_debt=0,
      start_date="",
      end_date="",
  )


@app.route("/logout")
def logout():
  session.clear()
  return redirect("/")


def run_bot():
    while True:
        try:
            bot.remove_webhook()
            bot.infinity_polling(timeout=60, long_polling_timeout=60)
        except Exception as e:
            print(f"Bot polling xatosi: {e}")


if __name__ == '__main__':
    init_web_db()

    bot_thread = threading.Thread(target=run_bot)
    bot_thread.daemon = True
    bot_thread.start()

    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
