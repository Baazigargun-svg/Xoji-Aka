from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def create_excel_invoice(order_id, shop_name, agent_name, date_str, price_type, cart_items, comment=''):
    wb = Workbook()
    ws = wb.active
    ws.title = f'Nakladnoy_{order_id}'
    ws.sheet_view.showGridLines = True
    
    title_font = Font(name='Arial', size=16, bold=True)
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    bold_font = Font(name='Arial', size=11, bold=True)
    header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
    total_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    
    ws.merge_cells('A1:E1')
    ws['A1'] = 'XOJI AKA FACTORY — NAKLADNOY'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')
    
    ws['A3'] = f'Buyurtma ID: #{order_id}'
    ws['A3'].font = bold_font
    ws['D3'] = f'Sana: {date_str}'
    ws['A4'] = f"Do'kon (Klient): {shop_name}"
    ws['D4'] = f'Narx turi: {price_type.upper()}'
    ws['A5'] = f'Agent: {agent_name}'
    if comment:
        ws['A6'] = f'Izoh (Kommentariya): {comment}'
        ws['A6'].font = bold_font

    start_row = 8 if comment else 7

    headers = ['№', 'Mahsulot nomi', "Miqdori", "Narxi (so'm)", 'Jami summa']
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    row_num = start_row + 1
    total_sum = 0
    for idx, item in enumerate(cart_items, 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=item['name']).alignment = Alignment(horizontal='left')
        ws.cell(row=row_num, column=3, value=item['qty']).alignment = Alignment(horizontal='right')
        ws.cell(row=row_num, column=4, value=item['price']).alignment = Alignment(horizontal='right')
        summa = item['qty'] * item['price']
        total_sum += summa
        ws.cell(row=row_num, column=5, value=summa).alignment = Alignment(horizontal='right')
        ws.cell(row=row_num, column=4).number_format = '#,##0'
        ws.cell(row=row_num, column=5).number_format = '#,##0'
        for col in range(1, 6):
            ws.cell(row=row_num, column=col).border = thin_border
        row_num += 1

    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
    ws.cell(row=row_num, column=1, value="JAMI TO'LOV:").alignment = Alignment(horizontal='right')
    ws.cell(row=row_num, column=1).font = bold_font
    total_val = ws.cell(row=row_num, column=5, value=total_sum)
    total_val.font = bold_font
    total_val.number_format = '#,##0'
    for col in range(1, 6):
        ws.cell(row=row_num, column=col).fill = total_fill
        ws.cell(row=row_num, column=col).border = thin_border

    row_num += 2
    ws.cell(row=row_num, column=2, value='Qabul qildim: ____')
    ws.cell(row=row_num, column=4, value='Topshirdim: ____')

    file_name = f'Nakladnoy_{order_id}.xlsx'
    wb.save(file_name)
    return file_name

def create_svodka_excel(aggregated_data, file_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Svodka"
    
    ws.cell(row=1, column=1, value="Mahsulot nomi").font = Font(bold=True)
    ws.cell(row=1, column=2, value="Jami miqdori").font = Font(bold=True)
    
    row = 2
    for p_name, qty in aggregated_data.items():
        ws.cell(row=row, column=1, value=p_name)
        ws.cell(row=row, column=2, value=qty)
        row += 1
        
    wb.save(file_path)
    return file_path

def create_excel_sex_income(income_id, staff_name, date_str, cart_items):
    wb = Workbook()
    ws = wb.active
    ws.title = f'Sex_Kirim_{income_id}'
    ws.sheet_view.showGridLines = True
    
    title_font = Font(name='Arial', size=16, bold=True)
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    bold_font = Font(name='Arial', size=11, bold=True)
    header_fill = PatternFill(start_color='006100', end_color='006100', fill_type='solid')
    total_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    
    ws.merge_cells('A1:D1')
    ws['A1'] = '🏭 SEXGA MAHSULOT KIRIM (SKLAD)'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')
    
    ws['A3'] = f'Kirim ID: #{income_id}'
    ws['A3'].font = bold_font
    ws['C3'] = f'Sana: {date_str}'
    ws['A4'] = f'Mas’ul xodim: {staff_name}'
    
    headers = ['№', 'Mahsulot nomi', "Miqdori (Soni/Kg)"]
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    row_num = 7
    total_qty = 0
    for idx, item in enumerate(cart_items, 1):
        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=item['name']).alignment = Alignment(horizontal='left')
        ws.cell(row=row_num, column=3, value=item['qty']).alignment = Alignment(horizontal='right')
        total_qty += item['qty']
        ws.cell(row=row_num, column=3).number_format = '#,##0.##'
        for col in range(1, 4):
            ws.cell(row=row_num, column=col).border = thin_border
        row_num += 1

    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=2)
    ws.cell(row=row_num, column=1, value="JAMI QO'SHILGAN MIQDOR:").alignment = Alignment(horizontal='right')
    ws.cell(row=row_num, column=1).font = bold_font
    total_val = ws.cell(row=row_num, column=3, value=total_qty)
    total_val.font = bold_font
    total_val.number_format = '#,##0.##'
    for col in range(1, 4):
        ws.cell(row=row_num, column=col).fill = total_fill
        ws.cell(row=row_num, column=col).border = thin_border

    file_name = f'Sex_Kirim_{income_id}.xlsx'
    wb.save(file_name)
    return file_name
